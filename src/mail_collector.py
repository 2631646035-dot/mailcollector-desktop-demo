"""
MailCollector - Outlook / QQ / IMAP 邮件收集器
多邻国风格 | PySide6 | 方案记忆 | 多条件筛选 | 定时执行
"""

import sys, os, json, re, zipfile, shutil, uuid, ssl
from datetime import datetime, timezone, timedelta
from pathlib import Path

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QStackedWidget, QPushButton, QLabel, QLineEdit, QTextEdit, QComboBox,
    QCheckBox, QFileDialog, QProgressBar, QScrollArea,
    QFrame, QSizePolicy, QMessageBox, QListWidget, QListWidgetItem,
    QDateEdit, QRadioButton, QButtonGroup, QDialog,
    QDialogButtonBox, QGraphicsDropShadowEffect, QTimeEdit, QSpinBox,
    QSystemTrayIcon, QMenu
)
from PySide6.QtCore import Qt, QDate, QThread, Signal, QSize, QTimer, QTime, QEvent
from PySide6.QtGui import QFont, QColor, QIcon

try:
    import win32com.client
except Exception:
    win32com = None

try:
    import keyring
except Exception:
    keyring = None

try:
    import win32cred
except Exception:
    win32cred = None


# ─── 配置 ───
APP_DIR = Path(os.path.expandvars(r"%APPDATA%")) / "MailCollector"
APP_DIR.mkdir(parents=True, exist_ok=True)
PLANS_FILE = APP_DIR / "plans.json"
APP_SETTINGS_FILE = APP_DIR / "settings.json"
DEFAULT_OUTPUT = Path.home() / "Desktop" / "MailCollector_Output"
CREDENTIAL_SERVICE = "MailCollector.IMAP"


IMAP_PRESETS = {
    "qq": {
        "name": "QQ 邮箱",
        "server": "imap.qq.com",
        "port": 993,
        "account_hint": "123456@qq.com",
        "password_hint": "请填写 QQ 邮箱授权码（不是 QQ 密码）",
        "help": "先在 QQ 邮箱网页版的「设置 > 账户」中开启 IMAP/SMTP，然后生成授权码。",
        "password_hint_en": "Enter the QQ Mail app password (not your QQ password)",
        "help_en": "Enable IMAP/SMTP in QQ Mail web Settings > Account, then generate an app password.",
    },
    "163": {
        "name": "163 邮箱",
        "server": "imap.163.com",
        "port": 993,
        "account_hint": "name@163.com",
        "password_hint": "请填写客户端授权码",
        "help": "请先在 163 邮箱网页版开启 IMAP/SMTP 服务并生成客户端授权码。",
        "password_hint_en": "Enter the client app password",
        "help_en": "Enable IMAP/SMTP in 163 Mail web settings and generate a client app password.",
    },
    "gmail": {
        "name": "Gmail",
        "server": "imap.gmail.com",
        "port": 993,
        "account_hint": "name@gmail.com",
        "password_hint": "请填写 Google 应用专用密码",
        "help": "开启两步验证后创建应用专用密码；普通 Google 密码通常不可用。",
        "password_hint_en": "Enter a Google app password",
        "help_en": "Enable 2-Step Verification, then create an app password. Your normal Google password usually will not work.",
    },
    "custom": {
        "name": "其他 IMAP",
        "server": "",
        "port": 993,
        "account_hint": "your@email.com",
        "password_hint": "请填写密码或应用授权码",
        "help": "请向邮箱服务商确认 IMAP 服务器、SSL 端口和应用授权码。",
        "password_hint_en": "Enter a password or app password",
        "help_en": "Confirm the IMAP server, SSL port and app-password requirements with your provider.",
    },
}


def credential_key(account):
    return (account or "").strip().lower()


def load_imap_secret(account):
    key = credential_key(account)
    if not key:
        return ""
    if win32cred:
        try:
            credential = win32cred.CredRead(
                f"{CREDENTIAL_SERVICE}:{key}",
                win32cred.CRED_TYPE_GENERIC,
            )
            secret = credential.get("CredentialBlob", "")
            return secret.decode("utf-16-le") if isinstance(secret, bytes) else str(secret)
        except Exception:
            pass
    if not keyring:
        return ""
    try:
        return keyring.get_password(CREDENTIAL_SERVICE, key) or ""
    except Exception:
        return ""


def save_imap_secret(account, secret):
    key = credential_key(account)
    if not key or not secret:
        return False
    if win32cred:
        try:
            win32cred.CredWrite({
                "Type": win32cred.CRED_TYPE_GENERIC,
                "TargetName": f"{CREDENTIAL_SERVICE}:{key}",
                "UserName": key,
                "CredentialBlob": secret,
                "Persist": win32cred.CRED_PERSIST_LOCAL_MACHINE,
            }, 0)
            return True
        except Exception:
            pass
    if not keyring:
        return False
    try:
        keyring.set_password(CREDENTIAL_SERVICE, key, secret)
        return True
    except Exception:
        return False


# ─── 主题、语言与全局样式 ───
THEMES = {
    "fresh": {
        "name": "清新果园",
        "primary": "#58CC02", "primary_dark": "#46A302",
        "bg": "#FFFFFF", "card": "#F7F7F7", "surface": "#FFFFFF",
        "text": "#3C3C3C", "secondary": "#777777", "border": "#E5E5E5",
        "blue": "#1CB0F6", "orange": "#FF9600", "red": "#FF4B4B", "warning": "#FFC800",
        "hover": "#F1F8ED", "radius": 14, "card_radius": 24,
        "font": '"Microsoft YaHei UI", "Segoe UI", sans-serif',
    },
    "pixel_farm": {
        "name": "像素田园",
        "primary": "#5C8A45", "primary_dark": "#365B32",
        "bg": "#F4E4B8", "card": "#FFF2C9", "surface": "#FFF9E7",
        "text": "#493526", "secondary": "#765A43", "border": "#B9824A",
        "blue": "#4C8391", "orange": "#C96F36", "red": "#A94438", "warning": "#E7B84A",
        "hover": "#EAD49A", "radius": 6, "card_radius": 10,
        "font": '"Cascadia Mono", "Microsoft YaHei UI", "Segoe UI", sans-serif',
    },
    "midnight": {
        "name": "深海夜色",
        "primary": "#69D2A0", "primary_dark": "#2D8C69",
        "bg": "#111827", "card": "#1E293B", "surface": "#172033",
        "text": "#F1F5F9", "secondary": "#A8B3C7", "border": "#34445E",
        "blue": "#60A5FA", "orange": "#F59E67", "red": "#FB7185", "warning": "#FACC15",
        "hover": "#25334A", "radius": 12, "card_radius": 20,
        "font": '"Microsoft YaHei UI", "Segoe UI", sans-serif',
    },
}


def load_app_settings():
    defaults = {"theme": "pixel_farm", "language": "zh_CN"}
    try:
        if APP_SETTINGS_FILE.exists():
            data = json.loads(APP_SETTINGS_FILE.read_text(encoding="utf-8"))
            defaults.update({k: v for k, v in data.items() if k in defaults})
    except Exception:
        pass
    if defaults["theme"] not in THEMES:
        defaults["theme"] = "pixel_farm"
    if defaults["language"] not in ("zh_CN", "en_US"):
        defaults["language"] = "zh_CN"
    return defaults


def save_app_settings(settings):
    APP_SETTINGS_FILE.write_text(
        json.dumps(settings, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


APP_SETTINGS = load_app_settings()
CURRENT_LANGUAGE = APP_SETTINGS["language"]


TRANSLATIONS = {
    "MailCollector - 邮件附件收集器": "MailCollector - Email & Attachment Collector",
    "安全邮件工作台": "Secure mail workspace", "收集助手": "Collector", "我的方案": "My Plans", "设置": "Settings",
    "最小化": "Minimize", "最大化": "Maximize", "还原": "Restore", "关闭": "Close",
    "⚙️\n设置": "⚙️\nSettings",
    "邮件收集助手": "Mail Collection Studio",
    "从 Outlook 自动收集邮件、附件和压缩包": "Collect emails, attachments and archives from Outlook",
    "通过官方 IMAP 安全收集 QQ 邮箱邮件与附件": "Securely collect QQ Mail messages and attachments via IMAP",
    "从支持 IMAP 的邮箱自动收集邮件与附件": "Collect messages and attachments from any IMAP mailbox",
    "邮箱来源": "Mailbox source", "Outlook 客户端": "Outlook desktop", "QQ 邮箱": "QQ Mail",
    "其他邮箱": "Other mailbox", "邮箱服务商": "Mail provider", "163 邮箱": "163 Mail",
    "自定义 IMAP": "Custom IMAP", "IMAP 服务器": "IMAP server", "端口": "Port",
    "邮箱账号": "Email address", "密码/授权码": "Password / app password", "显示授权码": "Show app password",
    "例如: imap.gmail.com": "For example: imap.gmail.com", "请输入密码或应用授权码": "Enter a password or app password",
    "🔌 测试连接": "🔌 Test connection", "正在连接...": "Connecting...",
    "1 选邮件": "1 Choose mail", "2 选文件": "2 Choose files", "3 自动收集": "3 Collect",
    "邮件筛选条件": "Email filters", "邮件标题": "Subject", "邮件正文关键词": "Body keywords",
    "发件人包含": "Sender contains", "收件人包含": "Recipient contains", "邮件日期范围": "Date range",
    "仅筛选有附件的邮件": "Only messages with attachments", "不限": "Any time", "当天": "Today",
    "近 7 天": "Last 7 days", "近 30 天": "Last 30 days", "近 3 个月": "Last 3 months", "今年": "This year", "自定义": "Custom",
    "从": "From", "到": "To", "标题包含": "Contains", "标题开头": "Starts with", "包含文字": "Contains", "以文字开头": "Starts with",
    "包含任意关键词": "Contains any keyword", "包含全部关键词": "Contains all keywords", "任一关键词": "Any keyword", "全部关键词": "All keywords",
    "例如：Invoice、账单、报价单": "For example: Invoice, bill, quotation",
    "多个关键词用英文逗号分隔，例如：invoice, payment, 账单": "Separate keywords with commas, e.g. invoice, payment, bill",
    "例如：@supplier.example": "For example: @supplier.example", "例如：@company.example": "For example: @company.example",
    "收集内容与保存方式": "Content & output", "你想收集什么": "What to collect",
    "下载附件": "Download attachments", "导出整封邮件（.msg / .eml）": "Export full message (.msg / .eml)",
    "附件和邮件都要": "Attachments and messages", "附件类型": "Attachment types", "自定义后缀": "Custom extensions",
    "自动解压压缩包 .zip / .rar / .7z": "Automatically extract .zip / .rar / .7z",
    ".msg, .txt，英文逗号分隔；留空表示不限": ".msg, .txt; comma-separated. Leave blank for all types",
    "保存位置": "Output folder", "📁 选择": "📁 Browse", "💾 保存为方案": "💾 Save plan", "🚀 立即执行": "🚀 Run now",
    "保存常用筛选规则，下次一键执行，也可以设置定时运行": "Save reusable filters, run them in one click, or schedule them",
    "🚀 执行选中方案": "🚀 Run selected", "🗑 删除选中": "🗑 Delete selected", "还没有保存方案": "No saved plans yet",
    "在「邮件收集助手」页面设置筛选条件后，可以保存为方案。": "Configure filters in Collector, then save them as a plan.",
    "明细": "Details", "编辑": "Edit", "⏰ 设置定时": "⏰ Schedule",
    "保存这个方案": "Save this plan", "可以覆盖原方案，也可以另存为一个新方案": "Overwrite the plan or save a new copy",
    "例如：法国一件代发": "For example: France dropshipping",
    "覆盖原方案": "Overwrite plan", "另存为新方案": "Save as new plan", "方案明细": "Plan details",
    "这个方案会按照下面的规则执行": "This plan will run with the rules below",
    "设置定时执行": "Schedule plan", "程序保持打开时，会按指定时间自动执行": "Runs automatically while the app remains open",
    "  ✅ 启用定时执行": "  ✅ Enable scheduled run", "点击上方绿色框即可开启或关闭定时执行": "Toggle scheduled runs using the control above",
    "每天执行时间": "Daily run time", "运行多久": "Run period", "往期文件处理方式": "Previous files",
    "每次执行都创建一个日期文件夹": "Create a dated folder for every run",
    "所有往期文件统一放在「往期」文件夹中": "Move all previous files into one archive folder",
    "外观与语言": "Appearance & language", "让 MailCollector 更像你自己的工作台": "Make MailCollector feel like your own workspace",
    "界面主题": "Interface theme", "界面语言": "Interface language", "清新果园": "Fresh Orchard",
    "像素田园": "Pixel Farm", "深海夜色": "Midnight Ocean", "简体中文": "Simplified Chinese", "English": "English",
    "主题和语言会立即生效，并在下次启动时保留。": "Theme and language apply immediately and persist for the next launch.",
    "像素田园是为 MailCollector 原创设计的温暖农场风格，不包含任何游戏素材。": "Pixel Farm is an original warm farm-inspired theme for MailCollector and contains no game assets.",
    "提示": "Notice", "错误": "Error", "配置有误": "Invalid configuration", "信息不完整": "Missing information",
    "任务完成": "Task complete", "确认删除": "Confirm deletion", "定时任务失败": "Scheduled task failed",
    "已保存": "Saved", "已覆盖": "Updated", "已进入编辑": "Edit mode", "定时已启用": "Schedule enabled", "定时已关闭": "Schedule disabled",
    "IMAP 连接正常": "IMAP connected", "IMAP 连接失败": "IMAP connection failed", "授权码未保存": "App password not saved",
}


def translate_text(text):
    if CURRENT_LANGUAGE != "en_US" or not isinstance(text, str) or not text:
        return text
    if text in TRANSLATIONS:
        return TRANSLATIONS[text]
    result = text
    for source, translated in sorted(TRANSLATIONS.items(), key=lambda item: len(item[0]), reverse=True):
        if len(source) >= 4:
            result = result.replace(source, translated)
    phrase_map = {
        "方案「": 'Plan "', "执行失败：": " failed:\n", "删除方案": "Delete plan ",
        "每天 ": "Daily at ", " 执行，至 ": ", until ", "未定时": "Not scheduled",
        "处理方式：": "Action: ", "创建：": "Created: ", "定时：": "Schedule: ",
        "定时执行：": "Scheduled run: ", " 个月": " months", "未设置": "Not set", "未执行": "Never",
        "━━━ 搜索条件 ━━━": "--- Search filters ---", "━━━ 扫描结果 ━━━": "--- Results ---",
        "━━━ 不匹配原因统计 ━━━": "--- Non-match reasons ---", "标题:": "Subject:", "正文:": "Body:",
        "发件人:": "Sender:", "收件人:": "Recipient:", "日期从:": "Date from:", "日期到:": "Date to:",
        "仅附件:": "Attachments only:", "附件类型:": "Attachment types:", "自动解压:": "Auto extract:",
        "扫描邮件:": "Scanned:", "匹配邮件:": "Matched:", "保存文件:": "Saved files:", "跳过重复:": "Duplicates skipped:",
        "保存位置:": "Output:", "封邮件": " messages", " 封": " messages", " 个": " items",
        "🔍 正在连接 IMAP 服务器...": "Connecting to the IMAP server...", "🔍 正在搜索邮件...": "Searching mail...",
        "📬 待检查邮件:": "📬 Messages to inspect:", "（从最新邮件开始）": " (newest first)",
        "✅ 匹配:": "✅ Match:", "📎 保存:": "📎 Saved:", "✉ 导出邮件:": "✉ Exported:", "↩ 跳过重复": "↩ Duplicate skipped",
        "执行失败:\n": "Run failed:\n", "请先选择一个方案": "Select a plan first",
        "请选择保存目录": "Choose an output folder", "请输入方案名称": "Enter a plan name",
        "━━━ 基本信息 ━━━": "--- Basic information ---", "━━━ 邮件筛选 ━━━": "--- Email filters ---",
        "━━━ 附件处理 ━━━": "--- Attachments ---", "━━━ 定时执行 ━━━": "--- Schedule ---",
        "邮箱来源：": "Mailbox source: ", "邮箱账号：": "Email address: ", "处理方式：": "Action: ", "保存位置：": "Output folder: ",
        "创建时间：": "Created: ", "标题规则：": "Subject rule: ", "标题文字：": "Subject text: ",
        "正文规则：": "Body rule: ", "正文关键词：": "Body keywords: ", "发件人包含：": "Sender contains: ", "收件人包含：": "Recipient contains: ",
        "日期范围：": "Date range: ", "仅有附件：": "Has attachments only: ", "附件类型：": "Attachment types: ", "自动解压：": "Auto extract: ",
        "是否启用：": "Enabled: ", "执行时间：": "Run time: ", "运行期限：": "Run period: ", "结束日期：": "End date: ",
        "往期方式：": "Archive mode: ", "上次执行：": "Last run: ", "未设置": "Not set", "不限": "Any", "是\n": "Yes\n", "否\n": "No\n",
        "每次按日期建文件夹": "Create a dated folder for each run", "统一放入往期文件夹": "Use one archive folder",
        "说明：定时任务执行前，会先把输出目录中的旧文件移入「往期」。新收集到的文件会留在输出目录根目录。": "Before a scheduled run, old output files are moved into the archive. Newly collected files remain in the output root.",
    }
    for source, translated in phrase_map.items():
        result = result.replace(source, translated)
    if 'Plan "' in result:
        result = result.replace("」", '"')
    return result


def set_theme_constants(theme_name):
    global C_PRIMARY, C_PRIMARY_DARK, C_BG, C_CARD, C_SURFACE, C_TEXT
    global C_TEXT_SECONDARY, C_BORDER, C_ACCENT_BLUE, C_ACCENT_ORANGE, C_ACCENT_RED, C_WARNING
    palette = THEMES.get(theme_name, THEMES["pixel_farm"])
    C_PRIMARY = palette["primary"]
    C_PRIMARY_DARK = palette["primary_dark"]
    C_BG = palette["bg"]
    C_CARD = palette["card"]
    C_SURFACE = palette["surface"]
    C_TEXT = palette["text"]
    C_TEXT_SECONDARY = palette["secondary"]
    C_BORDER = palette["border"]
    C_ACCENT_BLUE = palette["blue"]
    C_ACCENT_ORANGE = palette["orange"]
    C_ACCENT_RED = palette["red"]
    C_WARNING = palette["warning"]


set_theme_constants(APP_SETTINGS["theme"])


def build_style(theme_name):
    p = THEMES.get(theme_name, THEMES["pixel_farm"])
    radius = p["radius"]
    pixel_border = 3 if theme_name == "pixel_farm" else 2
    return f"""
QMainWindow, QWidget {{ background-color: {p['bg']}; color: {p['text']}; font-family: {p['font']}; font-size: 14px; }}
QLabel {{ color: {p['text']}; background: transparent; border: none; }}
QDialog, QMessageBox {{ background: {p['card']}; color: {p['text']}; }}
QLineEdit, QTextEdit, QComboBox, QDateEdit, QTimeEdit, QSpinBox {{
    border: {pixel_border}px solid {p['border']}; border-radius: {radius}px; padding: 9px 13px;
    background: {p['surface']}; color: {p['text']}; selection-background-color: {p['primary']};
}}
QLineEdit:focus, QTextEdit:focus, QComboBox:focus, QDateEdit:focus, QTimeEdit:focus, QSpinBox:focus {{ border-color: {p['primary']}; }}
QLineEdit:disabled, QComboBox:disabled {{ background: {p['card']}; color: {p['secondary']}; }}
QPushButton {{ border: {pixel_border}px solid {p['border']}; border-radius: {radius}px; padding: 10px 20px; font-weight: 700; background: {p['surface']}; color: {p['text']}; }}
QPushButton:hover {{ background: {p['hover']}; border-color: {p['primary']}; }}
QPushButton:pressed {{ background: {p['primary']}; color: white; }}
QScrollArea {{ border: none; background: transparent; }}
QScrollBar:vertical {{ background: {p['card']}; width: 16px; margin: 3px; border: {pixel_border}px solid {p['border']}; border-radius: {radius}px; }}
QScrollBar::handle:vertical {{ background: {p['primary']}; min-height: 38px; border: 2px solid {p['primary_dark']}; border-radius: {max(3, radius - 2)}px; }}
QScrollBar::handle:vertical:hover {{ background: {p['orange']}; border-color: {p['text']}; }}
QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{ height: 0px; }}
QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical {{ background: transparent; }}
QScrollBar:horizontal {{ background: {p['card']}; height: 14px; margin: 3px; border: {pixel_border}px solid {p['border']}; border-radius: {radius}px; }}
QScrollBar::handle:horizontal {{ background: {p['primary']}; min-width: 38px; border-radius: {max(3, radius - 2)}px; }}
QScrollBar::add-line:horizontal, QScrollBar::sub-line:horizontal {{ width: 0px; }}
QProgressBar {{ border: {pixel_border}px solid {p['border']}; border-radius: {radius}px; height: 12px; background: {p['surface']}; text-align: center; }}
QProgressBar::chunk {{ background: {p['primary']}; border-radius: {max(2, radius - 3)}px; }}
QListWidget {{ border: none; background: transparent; outline: none; }}
QListWidget::item {{ border: none; padding: 0px; margin: 6px 0px; background: transparent; }}
QCheckBox, QRadioButton {{ font-size: 14px; spacing: 8px; background: transparent; color: {p['text']}; }}
QComboBox::drop-down {{ border: none; width: 24px; }}
QComboBox QAbstractItemView {{ background: {p['surface']}; color: {p['text']}; border: {pixel_border}px solid {p['border']}; selection-background-color: {p['primary']}; selection-color: white; padding: 4px; }}
QToolTip {{ background: {p['text']}; color: {p['surface']}; border: 2px solid {p['border']}; padding: 6px; }}
QMessageBox QLabel {{ color: {p['text']}; min-width: 280px; padding: 8px; }}
QMessageBox QPushButton, QDialogButtonBox QPushButton {{ min-width: 88px; min-height: 34px; background: {p['surface']}; color: {p['text']}; border: {pixel_border}px solid {p['border']}; }}
QMessageBox QPushButton:hover, QDialogButtonBox QPushButton:hover {{ background: {p['primary']}; color: white; border-color: {p['primary_dark']}; }}
"""


def retranslate_widget_tree(root):
    widgets = [root] + root.findChildren(QWidget)
    text_widgets = (QLabel, QPushButton, QCheckBox, QRadioButton)
    for widget in widgets:
        if isinstance(widget, text_widgets):
            source = widget.property("i18n_source_text")
            if source is None:
                source = widget.text()
                widget.setProperty("i18n_source_text", source)
            widget.setText(translate_text(source))

        if isinstance(widget, QLineEdit):
            source = widget.property("i18n_source_placeholder")
            if source is None:
                source = widget.placeholderText()
                widget.setProperty("i18n_source_placeholder", source)
            widget.setPlaceholderText(translate_text(source))

        if isinstance(widget, QComboBox):
            sources = widget.property("i18n_combo_sources")
            if sources is None:
                sources = [widget.itemText(i) for i in range(widget.count())]
                widget.setProperty("i18n_combo_sources", sources)
            for i, source in enumerate(sources):
                if i < widget.count():
                    widget.setItemText(i, translate_text(source))

        tooltip_source = widget.property("i18n_tooltip_source")
        if tooltip_source is None and widget.toolTip():
            tooltip_source = widget.toolTip()
            widget.setProperty("i18n_tooltip_source", tooltip_source)
        if tooltip_source:
            widget.setToolTip(translate_text(tooltip_source))

        if isinstance(widget, QTextEdit) and widget.property("translatable_content"):
            source = widget.property("i18n_document_source")
            if source is None:
                source = widget.toPlainText()
                widget.setProperty("i18n_document_source", source)
            widget.setPlainText(translate_text(source))

    title_source = root.property("i18n_window_title")
    if title_source is None and root.windowTitle():
        title_source = root.windowTitle()
        root.setProperty("i18n_window_title", title_source)
    if title_source:
        root.setWindowTitle(translate_text(title_source))


QtMessageBox = QMessageBox


class QMessageBox:
    """使原有调用全部走同一套主题弹窗，并保留 Qt 标准返回值。"""
    StandardButton = QtMessageBox.StandardButton

    @staticmethod
    def _show(parent, icon, title, text, buttons=None, default_button=None):
        if buttons is None:
            buttons = QtMessageBox.StandardButton.Ok

        dialog = QDialog(parent)
        dialog.setModal(True)
        dialog.setWindowTitle(translate_text(title))
        dialog.setMinimumWidth(460)
        dialog.setMaximumWidth(680)
        dialog.setStyleSheet(build_style(APP_SETTINGS["theme"]))
        dialog.setProperty("message_result", QtMessageBox.StandardButton.NoButton)

        outer = QVBoxLayout(dialog)
        outer.setContentsMargins(24, 22, 24, 20)
        outer.setSpacing(16)

        header = QHBoxLayout()
        icons = {
            QtMessageBox.Icon.Information: ("✨", C_ACCENT_BLUE),
            QtMessageBox.Icon.Warning: ("⚠", C_WARNING),
            QtMessageBox.Icon.Critical: ("✕", C_ACCENT_RED),
            QtMessageBox.Icon.Question: ("?", C_ACCENT_ORANGE),
        }
        symbol, color = icons.get(icon, ("•", C_PRIMARY))
        icon_label = QLabel(symbol)
        icon_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon_label.setFixedSize(48, 48)
        icon_label.setStyleSheet(
            f"background: {C_SURFACE}; color: {color}; border: 3px solid {color}; "
            f"border-radius: {current_theme()['radius']}px; font-size: 24px; font-weight: 900;"
        )
        header.addWidget(icon_label)
        title_label = QLabel(translate_text(title))
        title_label.setWordWrap(True)
        title_label.setFont(QFont("Microsoft YaHei UI", 16, QFont.Bold))
        title_label.setStyleSheet(f"color: {C_TEXT}; background: transparent;")
        header.addWidget(title_label, 1)
        outer.addLayout(header)

        translated_text = translate_text(text)
        if len(translated_text) > 520 or translated_text.count("\n") > 12:
            body = QTextEdit()
            body.setReadOnly(True)
            body.setPlainText(translated_text)
            body.setMinimumHeight(260)
            body.setMaximumHeight(420)
        else:
            body = QLabel(translated_text)
            body.setWordWrap(True)
            body.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
            body.setStyleSheet(
                f"background: {C_SURFACE}; color: {C_TEXT}; border: 2px solid {C_BORDER}; "
                f"border-radius: {current_theme()['radius']}px; padding: 14px;"
            )
        outer.addWidget(body)

        button_row = QHBoxLayout()
        button_row.addStretch()
        button_specs = [
            (QtMessageBox.StandardButton.Yes, "Yes" if CURRENT_LANGUAGE == "en_US" else "是", True),
            (QtMessageBox.StandardButton.No, "No" if CURRENT_LANGUAGE == "en_US" else "否", False),
            (QtMessageBox.StandardButton.Ok, "OK", True),
            (QtMessageBox.StandardButton.Cancel, "Cancel" if CURRENT_LANGUAGE == "en_US" else "取消", False),
            (QtMessageBox.StandardButton.Close, "Close" if CURRENT_LANGUAGE == "en_US" else "关闭", False),
        ]
        for standard_button, label, primary in button_specs:
            if not (buttons & standard_button):
                continue
            button = QPushButton(label)
            button.setMinimumSize(96, 42)
            if primary:
                button.setStyleSheet(
                    f"QPushButton {{ background: {C_PRIMARY}; color: white; border: 3px solid {C_PRIMARY_DARK}; "
                    f"border-radius: {current_theme()['radius']}px; font-weight: 800; }} "
                    f"QPushButton:hover {{ background: {C_ACCENT_ORANGE}; border-color: {C_TEXT}; }}"
                )
            button.clicked.connect(
                lambda checked=False, value=standard_button: (
                    dialog.setProperty("message_result", value), dialog.accept()
                )
            )
            if default_button == standard_button or (default_button is None and primary):
                button.setDefault(True)
            button_row.addWidget(button)
        outer.addLayout(button_row)

        dialog.exec()
        return dialog.property("message_result")

    @classmethod
    def information(cls, parent, title, text, buttons=None, defaultButton=None):
        return cls._show(parent, QtMessageBox.Icon.Information, title, text, buttons, defaultButton)

    @classmethod
    def warning(cls, parent, title, text, buttons=None, defaultButton=None):
        return cls._show(parent, QtMessageBox.Icon.Warning, title, text, buttons, defaultButton)

    @classmethod
    def critical(cls, parent, title, text, buttons=None, defaultButton=None):
        return cls._show(parent, QtMessageBox.Icon.Critical, title, text, buttons, defaultButton)

    @classmethod
    def question(cls, parent, title, text, buttons=None, defaultButton=None):
        return cls._show(parent, QtMessageBox.Icon.Question, title, text, buttons, defaultButton)


# ─── 通用 UI 组件 ───

def apply_shadow(widget, blur=22, y=6, color=QColor(0, 0, 0, 28)):
    shadow = QGraphicsDropShadowEffect(widget)
    shadow.setBlurRadius(blur)
    shadow.setOffset(0, y)
    shadow.setColor(color)
    widget.setGraphicsEffect(shadow)


def current_theme():
    return THEMES.get(APP_SETTINGS["theme"], THEMES["pixel_farm"])


class GreenButton(QPushButton):
    def __init__(self, text, parent=None):
        super().__init__(text, parent)
        self.setMinimumHeight(48)
        self.setStyleSheet(f"""
            QPushButton {{
                background: {C_PRIMARY};
                color: white;
                border: none;
                border-radius: {current_theme()['radius']}px;
                padding: 12px 28px;
                font-size: 15px;
                font-weight: 800;
                border-bottom: 5px solid {C_PRIMARY_DARK};
            }}
            QPushButton:hover {{
                background: {C_PRIMARY};
            }}
            QPushButton:pressed {{
                border-bottom: 2px solid {C_PRIMARY_DARK};
                padding-top: 15px;
                padding-bottom: 9px;
            }}
        """)


class WhiteButton(QPushButton):
    def __init__(self, text, parent=None):
        super().__init__(text, parent)
        self.setMinimumHeight(48)
        self.setStyleSheet(f"""
            QPushButton {{
                background: {C_SURFACE};
                color: {C_PRIMARY};
                border: 2px solid {C_BORDER};
                border-radius: {current_theme()['radius']}px;
                padding: 12px 28px;
                font-size: 15px;
                font-weight: 800;
                border-bottom: 5px solid #D6D6D6;
            }}
            QPushButton:hover {{
                background: {current_theme()['hover']};
                border-color: {C_PRIMARY};
            }}
            QPushButton:pressed {{
                border-bottom: 2px solid #D6D6D6;
                padding-top: 15px;
                padding-bottom: 9px;
            }}
        """)


class DangerButton(QPushButton):
    def __init__(self, text, parent=None):
        super().__init__(text, parent)
        self.setMinimumHeight(44)
        self.setStyleSheet(f"""
            QPushButton {{
                background: {C_SURFACE};
                color: {C_ACCENT_RED};
                border: 2px solid {C_BORDER};
                border-radius: {current_theme()['radius']}px;
                padding: 10px 22px;
                font-size: 13px;
                font-weight: 700;
                border-bottom: 4px solid #D6D6D6;
            }}
            QPushButton:hover {{
                background: #FFF0F0;
                border-color: {C_ACCENT_RED};
            }}
            QPushButton:pressed {{
                border-bottom: 2px solid #D6D6D6;
                padding-top: 12px;
                padding-bottom: 8px;
            }}
        """)


class Card(QFrame):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("DuolingoCard")
        border = f"3px solid {C_BORDER}" if APP_SETTINGS["theme"] == "pixel_farm" else "none"
        radius = current_theme()["card_radius"]
        self.setStyleSheet(f"""
            QFrame#DuolingoCard {{
                background: {C_CARD};
                border-radius: {radius}px;
                border: {border};
            }}
        """)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        if APP_SETTINGS["theme"] == "pixel_farm":
            shadow_color = QColor(C_PRIMARY_DARK)
            shadow_color.setAlpha(75)
            apply_shadow(self, blur=3, y=4, color=shadow_color)
        else:
            apply_shadow(self, blur=20, y=5, color=QColor(0, 0, 0, 22))


class TitleLabel(QLabel):
    def __init__(self, text, parent=None):
        super().__init__(text, parent)
        self.setFont(QFont("Microsoft YaHei UI", 24, QFont.Bold))
        self.setStyleSheet(f"""
            QLabel {{
                color: {C_TEXT};
                background: transparent;
                border: none;
            }}
        """)


class SubtitleLabel(QLabel):
    def __init__(self, text, parent=None):
        super().__init__(text, parent)
        self.setFont(QFont("Microsoft YaHei UI", 12))
        self.setStyleSheet(f"""
            QLabel {{
                color: {C_TEXT_SECONDARY};
                background: transparent;
                border: none;
            }}
        """)


class SectionLabel(QLabel):
    def __init__(self, text, parent=None):
        super().__init__(text, parent)
        self.setFont(QFont("Microsoft YaHei UI", 15, QFont.Bold))
        self.setStyleSheet(f"""
            QLabel {{
                color: {C_TEXT};
                background: transparent;
                border: none;
            }}
        """)


class FieldLabel(QLabel):
    def __init__(self, text, parent=None):
        super().__init__(text, parent)
        self.setFont(QFont("Microsoft YaHei UI", 10, QFont.Bold))
        self.setMinimumWidth(110)
        self.setStyleSheet(f"""
            QLabel {{
                color: {C_TEXT_SECONDARY};
                background: transparent;
                border: none;
                padding-left: 2px;
            }}
        """)


class StepBadge(QLabel):
    def __init__(self, text, color=C_PRIMARY, parent=None):
        super().__init__(text, parent)
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setFont(QFont("Microsoft YaHei UI", 10, QFont.Bold))
        self.setStyleSheet(f"""
            QLabel {{
                background: {C_SURFACE};
                color: {color};
                border: 2px solid {C_BORDER};
                border-radius: {current_theme()['radius']}px;
                padding: 7px 12px;
            }}
        """)


class SmallActionButton(QPushButton):
    def __init__(self, text, color, parent=None):
        super().__init__(text, parent)
        self.setFixedSize(58, 32)
        self.setFont(QFont("Microsoft YaHei UI", 10, QFont.Bold))
        self.setStyleSheet(f"""
            QPushButton {{
                background: {C_SURFACE};
                color: {color};
                border: 1px solid {color};
                border-radius: 10px;
                padding: 0px;
            }}
            QPushButton:hover {{
                background: {color};
                color: white;
            }}
        """)


class SegmentedControl(QWidget):
    value_changed = Signal(str)

    def __init__(self, items, default_value=None, parent=None):
        super().__init__(parent)

        self.group = QButtonGroup(self)
        self.group.setExclusive(True)

        layout = QHBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(8)

        self.buttons = []

        for value, text in items:
            btn = QPushButton(text)
            btn.setCheckable(True)
            btn.setProperty("value", value)
            btn.setMinimumHeight(42)
            btn.setStyleSheet(f"""
                QPushButton {{
                    background: {C_SURFACE};
                    color: {C_TEXT_SECONDARY};
                    border: 2px solid {C_BORDER};
                    border-radius: {current_theme()['radius']}px;
                    padding: 9px 16px;
                    font-size: 13px;
                    font-weight: 700;
                }}
                QPushButton:hover {{
                    border-color: {C_PRIMARY};
                    color: {C_PRIMARY};
                    background: {current_theme()['hover']};
                }}
                QPushButton:checked {{
                    background: {C_PRIMARY};
                    color: white;
                    border-color: {C_PRIMARY};
                    border-bottom: 4px solid {C_PRIMARY_DARK};
                }}
            """)
            self.group.addButton(btn)
            layout.addWidget(btn)
            self.buttons.append(btn)
            btn.clicked.connect(lambda checked, v=value: self.value_changed.emit(v))

            if default_value is not None and value == default_value:
                btn.setChecked(True)

        if self.buttons and not any(b.isChecked() for b in self.buttons):
            self.buttons[0].setChecked(True)

    def value(self):
        checked = self.group.checkedButton()
        if checked:
            return checked.property("value")
        return None

    def set_value(self, value):
        for btn in self.buttons:
            if btn.property("value") == value:
                btn.setChecked(True)
                return
                return


class DateRangePicker(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(10)

        self.group = QButtonGroup(self)
        self.group.setExclusive(True)

        preset_row = QHBoxLayout()
        preset_row.setSpacing(8)

        self.buttons = []
        presets = [
            ("any", "不限"),
            ("today", "当天"),
            ("7d", "近 7 天"),
            ("30d", "近 30 天"),
            ("3m", "近 3 个月"),
            ("year", "今年"),
            ("custom", "自定义"),
        ]

        for value, text in presets:
            btn = QPushButton(text)
            btn.setCheckable(True)
            btn.setProperty("value", value)
            btn.setMinimumHeight(38)
            btn.setStyleSheet(f"""
                QPushButton {{
                    background: {C_SURFACE};
                    color: {C_TEXT_SECONDARY};
                    border: 2px solid {C_BORDER};
                    border-radius: {current_theme()['radius']}px;
                    padding: 7px 14px;
                    font-size: 13px;
                    font-weight: 700;
                }}
                QPushButton:hover {{
                    color: {C_PRIMARY};
                    border-color: {C_PRIMARY};
                    background: {current_theme()['hover']};
                }}
                QPushButton:checked {{
                    background: {C_PRIMARY};
                    color: white;
                    border-color: {C_PRIMARY};
                }}
            """)
            self.group.addButton(btn)
            preset_row.addWidget(btn)
            self.buttons.append(btn)

        preset_row.addStretch()
        root.addLayout(preset_row)

        self.custom_box = QWidget()
        custom_row = QHBoxLayout(self.custom_box)
        custom_row.setContentsMargins(0, 0, 0, 0)
        custom_row.setSpacing(8)

        custom_row.addWidget(QLabel("从"))
        self.date_from = QDateEdit()
        self.date_from.setCalendarPopup(True)
        self.date_from.setDate(QDate.currentDate().addMonths(-3))
        self.date_from.setMinimumDate(QDate(2020, 1, 1))
        custom_row.addWidget(self.date_from)

        custom_row.addWidget(QLabel("到"))
        self.date_to = QDateEdit()
        self.date_to.setCalendarPopup(True)
        self.date_to.setDate(QDate.currentDate())
        self.date_to.setMinimumDate(QDate(2020, 1, 1))
        custom_row.addWidget(self.date_to)

        custom_row.addStretch()
        root.addWidget(self.custom_box)

        self.group.buttonClicked.connect(self._on_changed)
        self.set_value("3m")

    def _on_changed(self, *args):
        self.custom_box.setVisible(self.value() == "custom")

    def value(self):
        checked = self.group.checkedButton()
        if checked:
            return checked.property("value")
        return "3m"

    def set_value(self, value):
        for btn in self.buttons:
            if btn.property("value") == value:
                btn.setChecked(True)
                break
        self._on_changed()

    def set_range(self, date_from, date_to):
        if not date_from and not date_to:
            self.set_value("any")
            return

        self.set_value("custom")

        if date_from:
            self.date_from.setDate(QDate(date_from.year, date_from.month, date_from.day))

        if date_to:
            self.date_to.setDate(QDate(date_to.year, date_to.month, date_to.day))

    def get_range(self):
        mode = self.value()
        today = QDate.currentDate()

        if mode == "any":
            return None, None

        if mode == "today":
            start = today
            end = today
        elif mode == "7d":
            start = today.addDays(-7)
            end = today
        elif mode == "30d":
            start = today.addDays(-30)
            end = today
        elif mode == "3m":
            start = today.addMonths(-3)
            end = today
        elif mode == "year":
            start = QDate(today.year(), 1, 1)
            end = today
        else:
            start = self.date_from.date()
            end = self.date_to.date()
            if start > end:
                start, end = end, start

        date_from = datetime(start.year(), start.month(), start.day(), 0, 0, 0, tzinfo=timezone.utc)
        date_to = datetime(end.year(), end.month(), end.day(), 23, 59, 59, tzinfo=timezone.utc)

        return date_from, date_to


# ─── 工具函数 ───

def new_plan_id():
    return str(uuid.uuid4())


def action_to_text(action):
    return {
        "attachment": "下载附件",
        "email": "导出邮件",
        "both": "附件和邮件都要"
    }.get(action, action or "未知")


def title_mode_to_text(mode):
    return {
        "contains": "标题包含指定文字",
        "startswith": "标题以指定文字开头"
    }.get(mode, mode or "未设置")


def body_mode_to_text(mode):
    return {
        "any": "正文包含任一关键词",
        "all": "正文包含全部关键词"
    }.get(mode, mode or "未设置")


def date_mode_to_text(mode):
    return {
        "any": "不限",
        "today": "当天",
        "7d": "近 7 天",
        "30d": "近 30 天",
        "3m": "近 3 个月",
        "year": "今年",
        "custom": "自定义"
    }.get(mode, mode or "未设置")


def unique_path(path: Path):
    if not path.exists():
        return path

    stem = path.stem
    suffix = path.suffix
    parent = path.parent
    i = 1

    while True:
        candidate = parent / f"{stem}_{i}{suffix}"
        if not candidate.exists():
            return candidate
        i += 1


def safe_extract_zip(zip_path, output_dir):
    """拒绝会写到目标文件夹之外的 ZIP 成员。"""
    target = Path(output_dir).resolve()
    with zipfile.ZipFile(zip_path, "r") as archive:
        for member in archive.infolist():
            destination = (target / member.filename).resolve()
            try:
                destination.relative_to(target)
            except ValueError:
                raise RuntimeError(f"压缩包含有不安全路径：{member.filename}")
        archive.extractall(target)


def refresh_dynamic_dates(conditions):
    """
    如果保存的是“当天 / 近7天”等动态日期，每次执行前重新计算。
    """
    c = dict(conditions)
    mode = c.get("date_mode", "custom")

    today = QDate.currentDate()

    if mode == "any":
        c["date_from"] = None
        c["date_to"] = None
        return c

    if mode == "today":
        start = today
        end = today
    elif mode == "7d":
        start = today.addDays(-7)
        end = today
    elif mode == "30d":
        start = today.addDays(-30)
        end = today
    elif mode == "3m":
        start = today.addMonths(-3)
        end = today
    elif mode == "year":
        start = QDate(today.year(), 1, 1)
        end = today
    else:
        return c

    c["date_from"] = datetime(start.year(), start.month(), start.day(), 0, 0, 0, tzinfo=timezone.utc)
    c["date_to"] = datetime(end.year(), end.month(), end.day(), 23, 59, 59, tzinfo=timezone.utc)

    return c


def archive_current_files(output_dir, archive_mode, schedule=None):
    """
    定时执行前，把输出根目录中已有文件移动到"往期"。
    archive_mode:
      by_date: 每次创建日期文件夹
      common: 统一放到"往期"

    高频模式下，每天只归档一次（通过 schedule.last_archive_date 控制）。
    """
    today_str = datetime.now().strftime("%Y-%m-%d")
    if schedule:
        if schedule.get("last_archive_date") == today_str:
            return
        schedule["last_archive_date"] = today_str

    root = Path(output_dir)
    root.mkdir(parents=True, exist_ok=True)

    archive_root = root / "往期"
    archive_root.mkdir(exist_ok=True)

    if archive_mode == "by_date":
        target = archive_root / datetime.now().strftime("%Y-%m-%d_%H%M%S")
        target.mkdir(exist_ok=True)
    else:
        target = archive_root

    for child in list(root.iterdir()):
        if child.name == "往期":
            continue

        dest = unique_path(target / child.name)

        try:
            shutil.move(str(child), str(dest))
        except Exception:
            pass


# ─── Outlook 邮件搜索与处理线程 ───

class OutlookWorker(QThread):
    progress = Signal(str)
    finished = Signal(dict)
    error = Signal(str)

    def __init__(self, conditions, output_dir, action):
        super().__init__()
        self.conditions = conditions
        self.output_dir = output_dir
        self.action = action

    def run(self):
        try:
            if win32com is None:
                raise RuntimeError("未检测到 pywin32。请先安装：pip install pywin32")

            outlook = win32com.client.Dispatch("Outlook.Application")
            namespace = outlook.GetNamespace("MAPI")
            inbox = namespace.GetDefaultFolder(6)

            self.progress.emit("🔍 正在搜索邮件...")

            messages = inbox.Items
            messages.Sort("[ReceivedTime]", True)

            matched = 0
            saved = 0
            skipped_duplicate = 0
            errors = []
            saved_ids = []

            total = messages.Count
            self.progress.emit(f"📬 扫描 {total} 封邮件...")

            debug_count = 0
            all_mails = []

            for i in range(1, total + 1):
                try:
                    mail = messages.Item(i)
                except Exception:
                    continue

                try:
                    subject_preview = (mail.Subject or "")[:80]
                    sender_preview = mail.SenderEmailAddress or ""
                    date_preview = str(mail.ReceivedTime)[:19]
                    attachments_preview = [att.FileName[:80] for att in mail.Attachments]
                except Exception:
                    continue

                debug_count += 1
                ok, reason = self._match(mail)

                all_mails.append({
                    "subject": subject_preview,
                    "sender": sender_preview,
                    "date": date_preview,
                    "attachments": attachments_preview,
                    "matched": ok,
                    "reason": reason
                })

                if debug_count <= 30 and not ok:
                    self.progress.emit(f"✗ #{debug_count} {reason}: {subject_preview[:50]}")

                if not ok:
                    continue

                matched += 1
                try:
                    saved_ids.append(mail.EntryID)
                except Exception:
                    pass
                self.progress.emit(f"✅ 匹配: {subject_preview[:50]}")

                if self.action in ("attachment", "both"):
                    exts = self.conditions.get("attachment_exts", [])
                    name_keywords = self.conditions.get("attachment_name_text", "").strip()
                    auto_unzip = self.conditions.get("auto_unzip", True)
                    zip_exts = {".zip", ".rar", ".7z"}

                    for att in mail.Attachments:
                        fname = att.FileName
                        _, file_ext = os.path.splitext(fname.lower())

                        if exts and file_ext not in [e.lower() for e in exts]:
                            continue

                        if name_keywords:
                            kwlist = [k.strip().lower() for k in name_keywords.split(",") if k.strip()]
                            if not any(k in fname.lower() for k in kwlist):
                                continue

                        if self.conditions.get("_dedupe") and self._file_exists_anywhere(fname):
                            skipped_duplicate += 1
                            self.progress.emit(f"  ↩ 跳过重复附件: {fname}")
                            continue

                        save_path = os.path.join(self.output_dir, fname)
                        base, e = os.path.splitext(fname)
                        c = 1

                        while os.path.exists(save_path):
                            save_path = os.path.join(self.output_dir, f"{base}_{c}{e}")
                            c += 1

                        try:
                            att.SaveAsFile(save_path)
                            saved += 1
                            self.progress.emit(f"  📎 保存: {os.path.basename(save_path)}")

                            if auto_unzip and file_ext in zip_exts:
                                try:
                                    unzip_dir = os.path.join(self.output_dir, base)
                                    os.makedirs(unzip_dir, exist_ok=True)

                                    if file_ext == ".zip":
                                        safe_extract_zip(save_path, unzip_dir)
                                    elif file_ext == ".rar":
                                        OutlookWorker._extract_rar(save_path, unzip_dir)
                                    elif file_ext == ".7z":
                                        OutlookWorker._extract_7z(save_path, unzip_dir)

                                    self.progress.emit(f"  📂 已解压到: {os.path.basename(unzip_dir)}")

                                except Exception as ze:
                                    self.progress.emit(f"  ⚠ 解压失败: {ze}")

                        except Exception as e:
                            errors.append(f"{fname}: {e}")

                if self.action in ("email", "both"):
                    fname = self._sanitize_filename(mail.Subject or "未命名邮件") + ".msg"

                    if self.conditions.get("_dedupe") and self._file_exists_anywhere(fname):
                        skipped_duplicate += 1
                        self.progress.emit(f"  ↩ 跳过重复邮件: {fname}")
                        continue

                    save_path = os.path.join(self.output_dir, fname)
                    c = 1

                    while os.path.exists(save_path):
                        fname = self._sanitize_filename(mail.Subject or "未命名邮件") + f"_{c}.msg"
                        save_path = os.path.join(self.output_dir, fname)
                        c += 1

                    try:
                        mail.SaveAs(save_path)
                        saved += 1
                        self.progress.emit(f"  ✉ 导出邮件: {fname}")
                    except Exception as e:
                        errors.append(f"邮件 {fname}: {e}")

                if matched >= 500:
                    self.progress.emit("⚠ 已达 500 封上限，停止扫描")
                    break

            fail_reasons = {}
            for m in all_mails:
                if not m["matched"]:
                    r = m["reason"]
                    fail_reasons[r] = fail_reasons.get(r, 0) + 1

            result = {
                "scanned": total,
                "matched": matched,
                "saved": saved,
                "skipped_duplicate": skipped_duplicate,
                "errors": errors,
                "output": self.output_dir,
                "conditions": self.conditions,
                "all_mails": all_mails,
                "fail_reasons": fail_reasons,
                "saved_ids": saved_ids
            }

            self.finished.emit(result)

        except Exception as e:
            self.error.emit(str(e))

    def _file_exists_anywhere(self, filename):
        root = self.conditions.get("_dedupe_root")
        if not root:
            return False

        try:
            root_path = Path(root)
            if not root_path.exists():
                return False

            for p in root_path.rglob(filename):
                if p.is_file():
                    return True
        except Exception:
            return False

        return False

    def _normalize_dt(self, dt):
        try:
            if isinstance(dt, datetime):
                if dt.tzinfo is None:
                    return dt.replace(tzinfo=timezone.utc)
                return dt.astimezone(timezone.utc)
        except Exception:
            pass
        return dt

    def _match(self, mail):
        c = self.conditions

        try:
            subject = mail.Subject or ""
            body = mail.Body or ""
            sender = mail.SenderEmailAddress or ""
            recipients = ""

            try:
                recips = mail.Recipients
                recipients = ", ".join(str(r.Address) for r in recips if r.Address)
            except Exception:
                pass

            received = self._normalize_dt(mail.ReceivedTime)

        except Exception:
            return False, "parse_error"

        try:
            if c.get("date_from") and received < c["date_from"]:
                return False, "date_before_range"

            if c.get("date_to") and received > c["date_to"]:
                return False, "date_after_range"
        except Exception:
            return False, "date_compare_error"

        title_text = c.get("title_text", "").strip()
        if title_text:
            if c.get("title_mode") == "startswith":
                if not subject.lower().startswith(title_text.lower()):
                    return False, "title_startswith"
            else:
                if title_text.lower() not in subject.lower():
                    return False, "title_contains"

        body_text = c.get("body_text", "").strip()
        if body_text:
            keywords = [k.strip() for k in body_text.split(",") if k.strip()]
            if c.get("body_mode") == "any":
                if not any(k.lower() in body.lower() for k in keywords):
                    return False, "body_keyword"
            else:
                if not all(k.lower() in body.lower() for k in keywords):
                    return False, "body_keyword"

        sender_text = c.get("sender_text", "").strip()
        if sender_text:
            if sender_text.lower() not in sender.lower():
                return False, "sender"

        recip_text = c.get("recipient_text", "").strip()
        if recip_text:
            if recip_text.lower() not in recipients.lower():
                return False, "recipient"

        if c.get("has_attachments"):
            try:
                if mail.Attachments.Count == 0:
                    return False, "no_attachments"
            except Exception:
                return False, "attachments_error"

        exclude_ids = c.get("_exclude_ids")
        if exclude_ids:
            try:
                entry_id = mail.EntryID
                if entry_id in exclude_ids:
                    return False, "already_collected"
            except Exception:
                pass

        return True, ""

    def _sanitize_filename(self, s):
        return re.sub(r'[\\/*?:"<>|]', "_", s)[:100]

    @staticmethod
    def _extract_rar(rar_path, out_dir):
        import subprocess

        unrar_paths = [
            r"C:\Program Files\WinRAR\UnRAR.exe",
            r"C:\Program Files (x86)\WinRAR\UnRAR.exe",
        ]

        for unrar in unrar_paths:
            if os.path.exists(unrar):
                subprocess.run([unrar, "x", "-y", rar_path, out_dir], capture_output=True, check=True)
                return

        raise RuntimeError("未找到 WinRAR/UnRAR，请安装 WinRAR 或使用 .zip 格式")

    @staticmethod
    def _extract_7z(sz_path, out_dir):
        import subprocess

        sz_paths = [
            r"C:\Program Files\7-Zip\7z.exe",
            r"C:\Program Files (x86)\7-Zip\7z.exe",
        ]

        for sz in sz_paths:
            if os.path.exists(sz):
                subprocess.run([sz, "x", f"-o{out_dir}", "-y", sz_path], capture_output=True, check=True)
                return

        raise RuntimeError("未找到 7-Zip，请安装 7-Zip 或使用 .zip 格式")


# ─── IMAP Worker ───

def open_imap_connection(config):
    import imaplib

    server = config.get("server", "").strip()
    account = config.get("account", "").strip()
    password = config.get("password", "")
    try:
        port = int(config.get("port", 993))
    except (TypeError, ValueError):
        raise RuntimeError("IMAP 端口必须是数字")

    if not server or not account or not password:
        raise RuntimeError("请填写 IMAP 服务器、邮箱账号和授权码")

    try:
        conn = imaplib.IMAP4_SSL(
            server,
            port,
            ssl_context=ssl.create_default_context(),
            timeout=20,
        )
        conn.login(account, password)
        status, _ = conn.select("INBOX", readonly=True)
        if status != "OK":
            raise RuntimeError("无法打开收件箱")
        return conn
    except imaplib.IMAP4.error as exc:
        message = str(exc)
        if "auth" in message.lower() or "login" in message.lower():
            raise RuntimeError("登录失败：请确认已开启 IMAP 服务，并使用授权码（不是邮箱登录密码）")
        raise RuntimeError(f"IMAP 服务器拒绝连接：{message}")
    except (OSError, TimeoutError) as exc:
        raise RuntimeError(f"无法连接 {server}:{port}，请检查网络、服务器和端口：{exc}")


class IMAPConnectionWorker(QThread):
    succeeded = Signal(str)
    error = Signal(str)

    def __init__(self, imap_config):
        super().__init__()
        self.imap_config = imap_config

    def run(self):
        conn = None
        try:
            conn = open_imap_connection(self.imap_config)
            status, data = conn.search(None, "ALL")
            count = len(data[0].split()) if status == "OK" and data else 0
            self.succeeded.emit(f"连接成功，收件箱共 {count} 封邮件。")
        except Exception as exc:
            self.error.emit(str(exc))
        finally:
            if conn:
                try:
                    conn.close()
                except Exception:
                    pass
                try:
                    conn.logout()
                except Exception:
                    pass

class IMAPWorker(QThread):
    progress = Signal(str)
    finished = Signal(dict)
    error = Signal(str)

    def __init__(self, conditions, output_dir, action, imap_config):
        super().__init__()
        self.conditions = conditions
        self.output_dir = output_dir
        self.action = action
        self.imap_config = imap_config

    def run(self):
        import email
        conn = None
        try:
            self.progress.emit("🔍 正在连接 IMAP 服务器...")
            conn = open_imap_connection(self.imap_config)

            self.progress.emit("🔍 正在搜索邮件...")

            search_args = ["ALL"]
            date_from = self.conditions.get("date_from")
            if date_from:
                search_args = ["SINCE", date_from.strftime("%d-%b-%Y")]
            status, data = conn.search(None, *search_args)
            if status != "OK":
                raise RuntimeError("服务器搜索邮件失败")
            mail_ids = list(reversed(data[0].split()))
            total = len(mail_ids)

            self.progress.emit(f"📬 待检查邮件: {total} 封（从最新邮件开始）")

            matched = 0
            saved = 0
            skipped_duplicate = 0
            errors = []
            saved_ids = []
            all_mails = []

            for mid in mail_ids:
                try:
                    fetch_status, msg_data = conn.fetch(mid, "(RFC822)")
                    if fetch_status != "OK" or not msg_data or not isinstance(msg_data[0], tuple):
                        continue
                    raw_email = msg_data[0][1]
                    msg = email.message_from_bytes(raw_email)
                except Exception:
                    continue

                subject = self._decode_header(msg.get("Subject", ""))
                sender = self._decode_header(msg.get("From", ""))
                recipients = self._decode_header(", ".join(filter(None, [msg.get("To", ""), msg.get("Cc", "")])))
                date_str = msg.get("Date", "")
                body = self._get_body(msg)

                has_att = False
                for part in msg.walk():
                    if part.get_content_disposition() == "attachment" or part.get_filename():
                        has_att = True
                        break

                received = None
                if date_str:
                    try:
                        from email.utils import parsedate_to_datetime
                        received = parsedate_to_datetime(date_str)
                        if received.tzinfo is None:
                            received = received.replace(tzinfo=timezone.utc)
                    except Exception:
                        pass

                ok, reason = self._match_imap(subject, body, sender, recipients, received, has_att)

                if ok:
                    exclude_ids = self.conditions.get("_exclude_ids")
                    if exclude_ids:
                        msg_id = msg.get("Message-ID", "").strip()
                        if msg_id and msg_id in exclude_ids:
                            ok = False
                            reason = "already_collected"

                all_mails.append({
                    "subject": subject[:80],
                    "sender": sender,
                    "date": str(received)[:19] if received else "",
                    "matched": ok,
                    "reason": reason
                })

                if not ok:
                    if len(all_mails) <= 30 and not ok:
                        self.progress.emit(f"✗ #{len(all_mails)} {reason}: {subject[:50]}")
                    continue

                matched += 1
                msg_id = msg.get("Message-ID", "").strip()
                if msg_id:
                    saved_ids.append(msg_id)
                self.progress.emit(f"✅ 匹配: {subject[:50]}")

                if self.action in ("attachment", "both"):
                    exts = self.conditions.get("attachment_exts", [])
                    name_keywords = self.conditions.get("attachment_name_text", "").strip()
                    auto_unzip = self.conditions.get("auto_unzip", True)
                    zip_exts = {".zip", ".rar", ".7z"}

                    for part in msg.walk():
                        if part.get_content_disposition() != "attachment" and not part.get_filename():
                            continue

                        fname = part.get_filename()
                        if not fname:
                            continue
                        fname = self._sanitize_filename(self._decode_header(fname))

                        _, file_ext = os.path.splitext(fname.lower())
                        if exts and file_ext not in [e.lower() for e in exts]:
                            continue

                        if name_keywords:
                            kwlist = [k.strip().lower() for k in name_keywords.split(",") if k.strip()]
                            if not any(k in fname.lower() for k in kwlist):
                                continue

                        if self.conditions.get("_dedupe") and self._file_exists_anywhere(fname):
                            skipped_duplicate += 1
                            self.progress.emit(f"  ↩ 跳过重复附件: {fname}")
                            continue

                        save_path = os.path.join(self.output_dir, fname)
                        base, e = os.path.splitext(fname)
                        c = 1
                        while os.path.exists(save_path):
                            save_path = os.path.join(self.output_dir, f"{base}_{c}{e}")
                            c += 1

                        try:
                            with open(save_path, "wb") as f:
                                f.write(part.get_payload(decode=True))
                            saved += 1
                            self.progress.emit(f"  📎 保存: {os.path.basename(save_path)}")

                            if auto_unzip and file_ext in zip_exts:
                                try:
                                    unzip_dir = os.path.join(self.output_dir, base)
                                    os.makedirs(unzip_dir, exist_ok=True)
                                    if file_ext == ".zip":
                                        safe_extract_zip(save_path, unzip_dir)
                                    elif file_ext == ".rar":
                                        OutlookWorker._extract_rar(save_path, unzip_dir)
                                    elif file_ext == ".7z":
                                        OutlookWorker._extract_7z(save_path, unzip_dir)
                                    self.progress.emit(f"  📂 已解压到: {os.path.basename(unzip_dir)}")
                                except Exception as ze:
                                    self.progress.emit(f"  ⚠ 解压失败: {ze}")
                        except Exception as ex:
                            errors.append(f"{fname}: {ex}")

                if self.action in ("email", "both"):
                    fname = self._sanitize_filename(subject or "未命名邮件") + ".eml"
                    if self.conditions.get("_dedupe") and self._file_exists_anywhere(fname):
                        skipped_duplicate += 1
                        self.progress.emit(f"  ↩ 跳过重复邮件: {fname}")
                    else:
                        save_path = unique_path(Path(self.output_dir) / fname)
                        try:
                            save_path.write_bytes(raw_email)
                            saved += 1
                            self.progress.emit(f"  ✉ 导出邮件: {save_path.name}")
                        except Exception as ex:
                            errors.append(f"邮件 {fname}: {ex}")

                if matched >= 500:
                    self.progress.emit("⚠ 已达 500 封匹配上限，停止扫描")
                    break

            fail_reasons = {}
            for m in all_mails:
                if not m["matched"]:
                    r = m["reason"]
                    fail_reasons[r] = fail_reasons.get(r, 0) + 1

            result = {
                "scanned": total,
                "matched": matched,
                "saved": saved,
                "skipped_duplicate": skipped_duplicate,
                "errors": errors,
                "output": self.output_dir,
                "conditions": self.conditions,
                "all_mails": all_mails,
                "fail_reasons": fail_reasons,
                "saved_ids": saved_ids
            }

            self.finished.emit(result)

        except Exception as e:
            self.error.emit(str(e))
        finally:
            if conn:
                try:
                    conn.close()
                except Exception:
                    pass
                try:
                    conn.logout()
                except Exception:
                    pass

    def _file_exists_anywhere(self, filename):
        root = self.conditions.get("_dedupe_root")
        if not root:
            return False
        try:
            root_path = Path(root)
            return root_path.exists() and any(p.is_file() for p in root_path.rglob(filename))
        except Exception:
            return False

    def _sanitize_filename(self, value):
        cleaned = re.sub(r'[\\/*?:"<>|]', "_", value or "")
        return cleaned.strip(" .")[:120] or "未命名"

    def _decode_header(self, h):
        from email.header import decode_header
        try:
            parts = decode_header(h)
            result = []
            for part, charset in parts:
                if isinstance(part, bytes):
                    result.append(part.decode(charset or "utf-8", errors="replace"))
                else:
                    result.append(str(part))
            return "".join(result)
        except:
            return str(h)

    def _get_body(self, msg):
        plain_parts = []
        html_parts = []
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_disposition() == "attachment":
                    continue
                content_type = part.get_content_type()
                if content_type in ("text/plain", "text/html"):
                    try:
                        charset = part.get_content_charset() or "utf-8"
                        value = part.get_payload(decode=True).decode(charset, errors="replace")
                        (plain_parts if content_type == "text/plain" else html_parts).append(value)
                    except Exception:
                        pass
        else:
            try:
                charset = msg.get_content_charset() or "utf-8"
                value = msg.get_payload(decode=True).decode(charset, errors="replace")
                if msg.get_content_type() == "text/html":
                    html_parts.append(value)
                else:
                    plain_parts.append(value)
            except Exception:
                plain_parts.append(str(msg.get_payload()))

        if plain_parts:
            return "\n".join(plain_parts)

        if html_parts:
            from html import unescape
            html_text = "\n".join(html_parts)
            html_text = re.sub(r"<(script|style)\b[^>]*>.*?</\1>", " ", html_text, flags=re.I | re.S)
            html_text = re.sub(r"<[^>]+>", " ", html_text)
            return re.sub(r"\s+", " ", unescape(html_text)).strip()

        return ""

    def _match_imap(self, subject, body, sender, recipients, received, has_att):
        c = self.conditions

        if c.get("date_from") and received:
            if received < c["date_from"]:
                return (False, "date_before_range")
        if c.get("date_to") and received:
            if received > c["date_to"]:
                return (False, "date_after_range")

        title_text = c.get("title_text", "").strip()
        if title_text:
            if c.get("title_mode") == "startswith":
                if not subject.lower().startswith(title_text.lower()):
                    return (False, "title_startswith")
            else:
                if title_text.lower() not in subject.lower():
                    return (False, "title_contains")

        body_text = c.get("body_text", "").strip()
        if body_text:
            keywords = [k.strip() for k in body_text.split(",") if k.strip()]
            if c.get("body_mode") == "any":
                if not any(k.lower() in body.lower() for k in keywords):
                    return (False, "body_keyword")
            else:
                if not all(k.lower() in body.lower() for k in keywords):
                    return (False, "body_keyword")

        sender_text = c.get("sender_text", "").strip()
        if sender_text:
            if sender_text.lower() not in sender.lower():
                return (False, "sender")

        recipient_text = c.get("recipient_text", "").strip()
        if recipient_text:
            if recipient_text.lower() not in recipients.lower():
                return (False, "recipient")

        if c.get("has_attachments") and not has_att:
            return (False, "no_attachments")

        return (True, "")


# ─── 方案管理 ───

def load_plans():
    changed = False

    if PLANS_FILE.exists():
        with open(PLANS_FILE, "r", encoding="utf-8") as f:
            plans = json.load(f)

        for plan in plans:
            if not plan.get("id"):
                plan["id"] = new_plan_id()
                changed = True

            if "schedule" not in plan:
                plan["schedule"] = {
                    "enabled": False,
                    "mode": "daily",
                    "time": "09:00",
                    "months": 1,
                    "end_date": "",
                    "archive_mode": "by_date",
                    "last_run": "",
                    "last_archive_date": "",
                }
                changed = True
            else:
                s = plan["schedule"]
                if "mode" not in s:
                    s["mode"] = "daily"
                    if "last_run_date" in s:
                        s["last_run"] = s.pop("last_run_date")
                    changed = True
                if "last_run" not in s and "last_run_date" in s:
                    s["last_run"] = s.pop("last_run_date")
                    changed = True
                if "last_archive_date" not in s:
                    s["last_archive_date"] = ""
                    changed = True
                if s.get("mode") == ScheduleDialog.MODE_INTERVAL:
                    if "window_start" not in s:
                        s["window_start"] = "12:00"
                        changed = True
                    if "window_end" not in s:
                        s["window_end"] = "18:00"
                        changed = True
                    if "pause_after_capture" not in s:
                        s["pause_after_capture"] = False
                        changed = True
                    if "collected_ids" not in s:
                        s["collected_ids"] = []
                        changed = True
                    if "paused_today" not in s:
                        s["paused_today"] = False
                        changed = True

            c = plan.get("conditions", {})

            imap_config = plan.get("imap_config") or {}
            legacy_secret = imap_config.pop("password", "")
            if legacy_secret:
                save_imap_secret(imap_config.get("account", ""), legacy_secret)
                changed = True

            if "date_mode" not in c:
                c["date_mode"] = "custom"
                changed = True

            for key in ("date_from", "date_to"):
                val = c.get(key)
                if val:
                    try:
                        c[key] = datetime.fromisoformat(val)
                        if c[key].tzinfo is None:
                            c[key] = c[key].replace(tzinfo=timezone.utc)
                        else:
                            c[key] = c[key].astimezone(timezone.utc)
                    except Exception:
                        c[key] = None

        if changed:
            save_plans(plans)

        return plans

    return []


def save_plans(plans):
    def convert(obj):
        if isinstance(obj, datetime):
            return obj.isoformat()
        if isinstance(obj, dict):
            # Credentials belong in Windows Credential Manager/keyring only.
            # Keeping this guard at the persistence boundary prevents a future
            # caller from accidentally writing an IMAP secret to plans.json.
            return {
                k: convert(v)
                for k, v in obj.items()
                if str(k).lower() != "password"
            }
        if isinstance(obj, list):
            return [convert(v) for v in obj]
        return obj

    with open(PLANS_FILE, "w", encoding="utf-8") as f:
        json.dump(convert(plans), f, ensure_ascii=False, indent=2)


# ─── 多邻国风格弹窗 ───

class StyledDialog(QDialog):
    def __init__(self, title="", fixed_size=None, parent=None):
        super().__init__(parent)
        self.setWindowTitle(title)

        if fixed_size:
            self.setFixedSize(*fixed_size)

        self.setStyleSheet(f"""
            QDialog {{
                background: {C_CARD};
            }}
            QLabel {{
                background: transparent;
                border: none;
            }}
            QRadioButton, QCheckBox {{
                font-size: 14px;
                font-weight: 700;
                color: {C_TEXT};
                background: transparent;
            }}
            QDialogButtonBox QPushButton {{
                background: {C_SURFACE};
                color: {C_PRIMARY};
                border: 2px solid {C_BORDER};
                border-radius: {current_theme()['radius']}px;
                padding: 8px 18px;
                font-size: 13px;
                font-weight: 800;
                min-width: 70px;
            }}
            QDialogButtonBox QPushButton:hover {{
                border-color: {C_PRIMARY};
                background: {current_theme()['hover']};
            }}
        """)

    def showEvent(self, event):
        retranslate_widget_tree(self)
        labels = {
            QDialogButtonBox.StandardButton.Ok: "OK",
            QDialogButtonBox.StandardButton.Cancel: "Cancel" if CURRENT_LANGUAGE == "en_US" else "取消",
            QDialogButtonBox.StandardButton.Close: "Close" if CURRENT_LANGUAGE == "en_US" else "关闭",
        }
        for button_box in self.findChildren(QDialogButtonBox):
            for standard_button, label in labels.items():
                button = button_box.button(standard_button)
                if button:
                    button.setText(label)
        super().showEvent(event)


class SavePlanDialog(StyledDialog):
    def __init__(self, default_name="", editing=False, parent=None):
        super().__init__("保存方案", fixed_size=(480, 250), parent=parent)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(26, 22, 26, 22)
        layout.setSpacing(14)

        head = QHBoxLayout()

        icon = QLabel("💾")
        icon.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon.setFixedSize(48, 48)
        icon.setStyleSheet("""
            QLabel {
                background: #E5F8D8;
                border-radius: 24px;
                font-size: 24px;
            }
        """)
        head.addWidget(icon)

        title_box = QVBoxLayout()
        title_box.setSpacing(2)
        title_box.addWidget(SectionLabel("保存这个方案"))
        title_box.addWidget(SubtitleLabel("可以覆盖原方案，也可以另存为一个新方案"))
        head.addLayout(title_box)
        head.addStretch()

        layout.addLayout(head)

        self.name_edit = QLineEdit()
        self.name_edit.setPlaceholderText("例如：法国一件代发")
        self.name_edit.setText(default_name)
        layout.addWidget(self.name_edit)

        self.overwrite_rb = QRadioButton("覆盖原方案")
        self.save_as_rb = QRadioButton("另存为新方案")

        self.mode_group = QButtonGroup(self)
        self.mode_group.addButton(self.overwrite_rb)
        self.mode_group.addButton(self.save_as_rb)

        if editing:
            self.overwrite_rb.setChecked(True)
        else:
            self.overwrite_rb.setEnabled(False)
            self.save_as_rb.setChecked(True)

        layout.addWidget(self.overwrite_rb)
        layout.addWidget(self.save_as_rb)

        btns = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        btns.accepted.connect(self.accept)
        btns.rejected.connect(self.reject)
        layout.addWidget(btns)

    def result_data(self):
        mode = "overwrite" if self.overwrite_rb.isChecked() else "save_as"
        return mode, self.name_edit.text().strip()


class PlanDetailDialog(StyledDialog):
    def __init__(self, plan, parent=None):
        super().__init__("方案明细", fixed_size=(620, 560), parent=parent)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(26, 22, 26, 22)
        layout.setSpacing(14)

        head = QHBoxLayout()

        icon = QLabel("📋")
        icon.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon.setFixedSize(50, 50)
        icon.setStyleSheet("""
            QLabel {
                background: #E8F5FF;
                border-radius: 25px;
                font-size: 24px;
            }
        """)
        head.addWidget(icon)

        title_box = QVBoxLayout()
        title_box.setSpacing(2)
        title_box.addWidget(SectionLabel(plan.get("name", "未命名方案")))
        title_box.addWidget(SubtitleLabel("这个方案会按照下面的规则执行"))
        head.addLayout(title_box)
        head.addStretch()

        layout.addLayout(head)

        card = QFrame()
        card.setStyleSheet(f"""
            QFrame {{
                background: {C_CARD};
                border-radius: 20px;
                border: none;
            }}
        """)
        apply_shadow(card, blur=16, y=4, color=QColor(0, 0, 0, 18))

        card_layout = QVBoxLayout(card)
        card_layout.setContentsMargins(18, 16, 18, 16)
        card_layout.setSpacing(10)

        text = QTextEdit()
        text.setReadOnly(True)
        text.setStyleSheet(f"""
            QTextEdit {{
                background: white;
                border: 2px solid {C_BORDER};
                border-radius: 16px;
                padding: 12px;
                font-size: 13px;
                color: {C_TEXT};
            }}
        """)

        c = plan.get("conditions", {})
        schedule = plan.get("schedule", {})

        detail = ""
        detail += "━━━ 基本信息 ━━━\n"
        source_names = {"outlook": "Outlook 客户端", "qq": "QQ 邮箱", "imap": "IMAP 邮箱"}
        detail += f"邮箱来源：{source_names.get(plan.get('mail_source', 'outlook'), '未知')}\n"
        if plan.get("mail_source") in ("qq", "imap"):
            imap_cfg = plan.get("imap_config") or {}
            detail += f"邮箱账号：{imap_cfg.get('account', '未设置')}\n"
        detail += f"处理方式：{action_to_text(plan.get('action'))}\n"
        detail += f"保存位置：{plan.get('output_dir', '')}\n"
        detail += f"创建时间：{plan.get('created', '')}\n\n"

        detail += "━━━ 邮件筛选 ━━━\n"
        detail += f"标题规则：{title_mode_to_text(c.get('title_mode'))}\n"
        detail += f"标题文字：{c.get('title_text') or '未设置'}\n"
        detail += f"正文规则：{body_mode_to_text(c.get('body_mode'))}\n"
        detail += f"正文关键词：{c.get('body_text') or '未设置'}\n"
        detail += f"发件人包含：{c.get('sender_text') or '未设置'}\n"
        detail += f"收件人包含：{c.get('recipient_text') or '未设置'}\n"
        detail += f"日期范围：{date_mode_to_text(c.get('date_mode', 'custom'))}\n"
        detail += f"仅有附件：{'是' if c.get('has_attachments') else '否'}\n\n"

        detail += "━━━ 附件处理 ━━━\n"
        detail += f"附件类型：{', '.join(c.get('attachment_exts', [])) or '不限'}\n"
        detail += f"附件名称包含：{c.get('attachment_name_text') or '不限'}\n"
        detail += f"自动解压：{'是' if c.get('auto_unzip') else '否'}\n\n"

        detail += "━━━ 定时执行 ━━━\n"
        detail += f"是否启用：{'是' if schedule.get('enabled') else '否'}\n"
        mode = schedule.get("mode", "daily")
        mode_names = {"daily": "每天一次", "interval": "间隔分钟", "weekly": "每周指定日"}
        detail += f"调度模式：{mode_names.get(mode, '每天一次')}\n"
        if mode == "daily":
            detail += f"执行时间：{schedule.get('time', '09:00')}\n"
        elif mode == "interval":
            detail += f"刷新间隔：每 {schedule.get('interval_minutes', 5)} 分钟\n"
        elif mode == "weekly":
            detail += f"执行时间：{schedule.get('weekly_time', '18:00')}\n"
            days = schedule.get("weekly_days", [])
            day_map = {1: "一", 2: "二", 3: "三", 4: "四", 5: "五", 6: "六", 7: "日"}
            day_str = ",".join(f"周{d}" for d in sorted(days)) if days else "未设置"
            detail += f"执行日：{day_str}\n"
        detail += f"运行期限：{schedule.get('months', 1)} 个月\n"
        detail += f"结束日期：{schedule.get('end_date') or '未设置'}\n"
        detail += f"往期方式：{'每次按日期建文件夹' if schedule.get('archive_mode') == 'by_date' else '统一放入往期文件夹'}\n"
        last_run = schedule.get("last_run", "")
        detail += f"上次执行：{last_run if last_run else '未执行'}\n"

        text.setText(detail)
        text.setProperty("translatable_content", True)
        card_layout.addWidget(text)

        layout.addWidget(card)

        btns = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
        btns.rejected.connect(self.reject)
        btns.accepted.connect(self.accept)
        layout.addWidget(btns)


class ScheduleDialog(StyledDialog):
    MODE_DAILY = "daily"
    MODE_INTERVAL = "interval"
    MODE_WEEKLY = "weekly"

    WEEKDAY_NAMES = {
        "zh_CN": ["周一", "周二", "周三", "周四", "周五", "周六", "周日"],
        "en_US": ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"],
    }

    def __init__(self, plan, parent=None):
        super().__init__("设置定时执行", fixed_size=(620, 610), parent=parent)

        icon_path = Path(__file__).resolve().parent / "mailcollector-icon.ico"
        if icon_path.exists():
            self.setWindowIcon(QIcon(str(icon_path)))

        self.plan = plan
        schedule = plan.get("schedule", {})

        layout = QVBoxLayout(self)
        layout.setContentsMargins(26, 22, 26, 22)
        layout.setSpacing(14)

        head = QHBoxLayout()

        icon = QLabel("⏰")
        icon.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon.setFixedSize(50, 50)
        icon.setStyleSheet("""
            QLabel {
                background: #FFF4D6;
                border-radius: 25px;
                font-size: 24px;
            }
        """)
        head.addWidget(icon)

        title_box = QVBoxLayout()
        title_box.setSpacing(2)
        title_box.addWidget(SectionLabel(f"定时执行：{plan.get('name', '')}"))
        title_box.addWidget(SubtitleLabel("程序保持打开时，会按指定时间自动执行"))
        head.addLayout(title_box)
        head.addStretch()

        layout.addLayout(head)

        card = QFrame()
        card.setStyleSheet(f"""
            QFrame {{
                background: {C_CARD};
                border-radius: 20px;
                border: none;
            }}
        """)
        apply_shadow(card, blur=16, y=4, color=QColor(0, 0, 0, 18))

        cl = QVBoxLayout(card)
        cl.setContentsMargins(18, 16, 18, 16)
        cl.setSpacing(12)

        self.enable_cb = QCheckBox("  ✅ 启用定时执行")
        self.enable_cb.setCursor(Qt.CursorShape.PointingHandCursor)
        self.enable_cb.setChecked(True)
        self.enable_cb.setMinimumHeight(58)
        self.enable_cb.setStyleSheet(f"""
            QCheckBox {{
                background: white;
                color: {C_TEXT};
                border: 2px solid {C_BORDER};
                border-radius: 18px;
                padding: 14px 18px;
                font-size: 15px;
                font-weight: 800;
            }}

            QCheckBox:hover {{
                border-color: {C_PRIMARY};
                background: #F8FFF3;
                color: {C_PRIMARY};
            }}

            QCheckBox:checked {{
                background: #E5F8D8;
                color: {C_PRIMARY_DARK};
                border: 2px solid {C_PRIMARY};
            }}

            QCheckBox::indicator {{
                width: 22px;
                height: 22px;
            }}
        """)

        cl.addWidget(self.enable_cb)

        tip_lbl = QLabel("点击上方绿色框即可开启或关闭定时执行")
        tip_lbl.setStyleSheet(f"""
            QLabel {{
                color: {C_TEXT_SECONDARY};
                font-size: 12px;
                padding-left: 4px;
            }}
        """)
        tip_lbl.setProperty("translatable_text", True)
        cl.addWidget(tip_lbl)

        cl.addWidget(FieldLabel("调度模式"))
        self.mode_combo = QComboBox()
        self.mode_combo.setMinimumWidth(280)
        self.mode_combo.addItem("🔄 间隔分钟（每N分钟执行一次）")
        self.mode_combo.addItem("📅 每天一次（指定时间执行）")
        self.mode_combo.addItem("📆 每周指定日（指定星期几执行）")
        self.mode_combo.setCurrentIndex(0)
        self.mode_combo.currentIndexChanged.connect(self._on_mode_changed)
        cl.addWidget(self.mode_combo)

        self.stack = QStackedWidget()
        self.stack.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.stack.setStyleSheet("background: transparent;")

        self.interval_page = QWidget()
        iv = QVBoxLayout(self.interval_page)
        iv.setContentsMargins(0, 8, 0, 0)
        iv.setSpacing(8)

        ir = QHBoxLayout()
        ir.addWidget(FieldLabel("刷新间隔"))
        self.interval_spin = QSpinBox()
        self.interval_spin.setMinimum(1)
        self.interval_spin.setMaximum(1440)
        self.interval_spin.setValue(int(schedule.get("interval_minutes", 10)))
        self.interval_spin.setSuffix(" 分钟")
        self.interval_spin.setFixedWidth(100)
        ir.addWidget(self.interval_spin)
        ir.addStretch()
        iv.addLayout(ir)

        wr = QHBoxLayout()
        wr.addWidget(FieldLabel("执行时间窗口"))
        wr.setSpacing(6)
        self.window_start = QTimeEdit()
        self.window_start.setDisplayFormat("HH:mm")
        self.window_start.setFixedWidth(90)
        try:
            h, m = schedule.get("window_start", "12:00").split(":")
            self.window_start.setTime(QTime(int(h), int(m)))
        except Exception:
            self.window_start.setTime(QTime(12, 0))
        wr.addWidget(self.window_start)
        wr.addWidget(QLabel("~"))
        self.window_end = QTimeEdit()
        self.window_end.setDisplayFormat("HH:mm")
        self.window_end.setFixedWidth(90)
        try:
            h, m = schedule.get("window_end", "18:00").split(":")
            self.window_end.setTime(QTime(int(h), int(m)))
        except Exception:
            self.window_end.setTime(QTime(18, 0))
        wr.addWidget(self.window_end)
        wr.addStretch()
        iv.addLayout(wr)

        self.pause_after_capture = QCheckBox("捕获到邮件后暂停当天执行，不再继续刷新")
        self.pause_after_capture.setChecked(schedule.get("pause_after_capture", False))
        self.pause_after_capture.setStyleSheet(f"color: {C_TEXT_SECONDARY}; font-size: 13px;")
        iv.addWidget(self.pause_after_capture)

        interval_hint = QLabel("在指定时间窗口内每N分钟检查一次，同一天的新邮件收集到同一个文件夹中。")
        interval_hint.setWordWrap(True)
        interval_hint.setStyleSheet(f"color: {C_TEXT_SECONDARY}; font-size: 12px;")
        interval_hint.setProperty("translatable_text", True)
        iv.addWidget(interval_hint)
        self.stack.addWidget(self.interval_page)

        self.daily_page = QWidget()
        dv = QVBoxLayout(self.daily_page)
        dv.setContentsMargins(0, 8, 0, 0)
        dv.setSpacing(8)
        dr = QHBoxLayout()
        dr.addWidget(FieldLabel("每天执行时间"))
        self.time_edit = QTimeEdit()
        self.time_edit.setDisplayFormat("HH:mm")
        self.time_edit.setFixedWidth(90)
        try:
            h, m = schedule.get("time", "09:00").split(":")
            self.time_edit.setTime(QTime(int(h), int(m)))
        except Exception:
            self.time_edit.setTime(QTime(9, 0))
        dr.addWidget(self.time_edit)
        dr.addStretch()
        dv.addLayout(dr)
        self.stack.addWidget(self.daily_page)

        self.weekly_page = QWidget()
        wv = QVBoxLayout(self.weekly_page)
        wv.setContentsMargins(0, 8, 0, 0)
        wv.setSpacing(10)
        wr = QHBoxLayout()
        wr.addWidget(FieldLabel("每周执行时间"))
        self.weekly_time_edit = QTimeEdit()
        self.weekly_time_edit.setDisplayFormat("HH:mm")
        self.weekly_time_edit.setFixedWidth(90)
        try:
            h, m = schedule.get("weekly_time", schedule.get("time", "18:00")).split(":")
            self.weekly_time_edit.setTime(QTime(int(h), int(m)))
        except Exception:
            self.weekly_time_edit.setTime(QTime(18, 0))
        wr.addWidget(self.weekly_time_edit)
        wr.addStretch()
        wv.addLayout(wr)
        wv.addWidget(FieldLabel("选择执行日（可多选）"))
        wdays_layout = QHBoxLayout()
        wdays_layout.setSpacing(6)
        self.weekday_checks = []
        saved_days = schedule.get("weekly_days", [])
        lang = CURRENT_LANGUAGE
        day_names = self.WEEKDAY_NAMES.get(lang, self.WEEKDAY_NAMES["zh_CN"])
        for i, name in enumerate(day_names):
            cb = QCheckBox(name)
            cb.setChecked((i + 1) in saved_days)
            cb.setStyleSheet(f"padding: 2px 4px; color: {C_TEXT};")
            self.weekday_checks.append(cb)
            wdays_layout.addWidget(cb)
        wdays_layout.addStretch()
        wv.addLayout(wdays_layout)
        weekly_hint = QLabel("选中多个则每周多天执行。只在勾选日期的指定时间执行一次。")
        weekly_hint.setWordWrap(True)
        weekly_hint.setStyleSheet(f"color: {C_TEXT_SECONDARY}; font-size: 12px;")
        weekly_hint.setProperty("translatable_text", True)
        wv.addWidget(weekly_hint)
        self.stack.addWidget(self.weekly_page)

        cl.addWidget(self.stack)

        row2 = QHBoxLayout()
        row2.addWidget(FieldLabel("运行多久"))
        self.month_spin = QSpinBox()
        self.month_spin.setMinimum(1)
        self.month_spin.setMaximum(6)
        self.month_spin.setValue(int(schedule.get("months", 1)))
        self.month_spin.setSuffix(" months" if CURRENT_LANGUAGE == "en_US" else " 个月")
        self.month_spin.setFixedWidth(100)
        row2.addWidget(self.month_spin)
        row2.addStretch()
        cl.addLayout(row2)

        cl.addWidget(FieldLabel("往期文件处理方式"))
        self.archive_group = QButtonGroup(self)
        self.archive_by_date = QRadioButton("每次执行都创建一个日期文件夹")
        self.archive_common = QRadioButton("所有往期文件统一放在「往期」文件夹中")
        self.archive_group.addButton(self.archive_by_date)
        self.archive_group.addButton(self.archive_common)
        if schedule.get("archive_mode", "by_date") == "common":
            self.archive_common.setChecked(True)
        else:
            self.archive_by_date.setChecked(True)
        cl.addWidget(self.archive_by_date)
        cl.addWidget(self.archive_common)

        hint = QLabel("说明：定时任务执行前，会先把输出目录中的旧文件移入「往期」。新收集到的文件会留在输出目录根目录。")
        hint.setWordWrap(True)
        hint.setStyleSheet(f"color: {C_TEXT_SECONDARY}; font-size: 12px;")
        hint.setProperty("translatable_text", True)
        cl.addWidget(hint)

        layout.addWidget(card)

        btns = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        btns.accepted.connect(self.accept)
        btns.rejected.connect(self.reject)
        layout.addWidget(btns)

        existing_mode = schedule.get("mode", self.MODE_DAILY)
        mode_map = {self.MODE_INTERVAL: 0, self.MODE_DAILY: 1, self.MODE_WEEKLY: 2}
        self.mode_combo.setCurrentIndex(mode_map.get(existing_mode, 1))
        self._on_mode_changed(self.mode_combo.currentIndex())

    def _on_mode_changed(self, index):
        self.stack.setCurrentIndex(index)

    def result_data(self):
        months = self.month_spin.value()
        end_date = (datetime.now() + timedelta(days=months * 30)).strftime("%Y-%m-%d")
        old_schedule = self.plan.get("schedule", {})

        mode_idx = self.mode_combo.currentIndex()
        mode_map = {0: self.MODE_INTERVAL, 1: self.MODE_DAILY, 2: self.MODE_WEEKLY}
        mode = mode_map.get(mode_idx, self.MODE_DAILY)

        result = {
            "enabled": self.enable_cb.isChecked(),
            "mode": mode,
            "months": months,
            "end_date": end_date,
            "archive_mode": "by_date" if self.archive_by_date.isChecked() else "common",
            "last_run": old_schedule.get("last_run", ""),
            "last_archive_date": old_schedule.get("last_archive_date", ""),
        }

        if mode == self.MODE_DAILY:
            result["time"] = self.time_edit.time().toString("HH:mm")
        elif mode == self.MODE_INTERVAL:
            result["interval_minutes"] = self.interval_spin.value()
            result["window_start"] = self.window_start.time().toString("HH:mm")
            result["window_end"] = self.window_end.time().toString("HH:mm")
            result["pause_after_capture"] = self.pause_after_capture.isChecked()
            result.setdefault("collected_ids", old_schedule.get("collected_ids", []))
            result.setdefault("paused_today", old_schedule.get("paused_today", False))
        elif mode == self.MODE_WEEKLY:
            result["weekly_time"] = self.weekly_time_edit.time().toString("HH:mm")
            days = []
            for i, cb in enumerate(self.weekday_checks):
                if cb.isChecked():
                    days.append(i + 1)
            result["weekly_days"] = days

        return result


# ─── 页面1: 收集助手 ───

class CreateTaskPage(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main = main_window

        self.editing_plan_id = None
        self.editing_plan_name = ""

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)

        container = QWidget()
        layout = QVBoxLayout(container)
        layout.setSpacing(18)
        layout.setContentsMargins(40, 30, 40, 30)

        header = QHBoxLayout()
        header.setSpacing(14)

        mascot = QLabel("🦉")
        mascot.setAlignment(Qt.AlignmentFlag.AlignCenter)
        mascot.setFixedSize(58, 58)
        mascot.setStyleSheet("""
            QLabel {
                background: #E5F8D8;
                border-radius: 29px;
                font-size: 30px;
                border: none;
            }
        """)
        header.addWidget(mascot)

        title_box = QVBoxLayout()
        title_box.setSpacing(4)
        title_box.addWidget(TitleLabel("邮件收集助手"))
        subtitle = SubtitleLabel("从 Outlook 自动收集邮件、附件和压缩包")
        subtitle.setObjectName("subtitle")
        title_box.addWidget(subtitle)
        header.addLayout(title_box)
        header.addStretch()

        layout.addLayout(header)

        # 邮箱类型选择
        source_card = Card()
        scl = QVBoxLayout(source_card)
        scl.setContentsMargins(24, 16, 24, 16)
        scl.setSpacing(10)
        scl.addWidget(FieldLabel("邮箱来源"))
        source_row = QHBoxLayout()
        source_row.setSpacing(12)
        self.mail_source = SegmentedControl([
            ("outlook", "Outlook 客户端"),
            ("qq", "QQ 邮箱"),
            ("imap", "其他邮箱"),
        ], default_value="outlook")
        self.mail_source.value_changed.connect(self._on_source_changed)
        source_row.addWidget(self.mail_source)
        source_row.addStretch()
        scl.addLayout(source_row)

        # IMAP 配置面板（默认隐藏）
        self.imap_panel = QWidget()
        self.imap_panel.setVisible(False)
        imap_layout = QVBoxLayout(self.imap_panel)
        imap_layout.setContentsMargins(0, 6, 0, 0)
        imap_layout.setSpacing(8)

        self.imap_provider_row = QWidget()
        provider_layout = QHBoxLayout(self.imap_provider_row)
        provider_layout.setContentsMargins(0, 0, 0, 0)
        provider_layout.setSpacing(10)
        provider_layout.addWidget(FieldLabel("邮箱服务商"))
        self.imap_provider = QComboBox()
        self.imap_provider.addItem("163 邮箱", "163")
        self.imap_provider.addItem("Gmail", "gmail")
        self.imap_provider.addItem("自定义 IMAP", "custom")
        self.imap_provider.currentIndexChanged.connect(self._on_provider_changed)
        provider_layout.addWidget(self.imap_provider)
        provider_layout.addStretch()
        imap_layout.addWidget(self.imap_provider_row)

        self.imap_help = QLabel("")
        self.imap_help.setWordWrap(True)
        self.imap_help.setStyleSheet(
            f"background: #FFF8DC; color: {C_TEXT}; border: 1px solid #F1D66A; "
            "border-radius: 12px; padding: 10px 12px;"
        )
        imap_layout.addWidget(self.imap_help)

        imap_row1 = QHBoxLayout()
        imap_row1.setSpacing(10)
        imap_left = QVBoxLayout()
        imap_left.setSpacing(4)
        imap_left.addWidget(FieldLabel("IMAP 服务器"))
        self.imap_server = QLineEdit()
        self.imap_server.setPlaceholderText("例如: imap.gmail.com")
        imap_left.addWidget(self.imap_server)
        imap_row1.addLayout(imap_left)
        imap_right = QVBoxLayout()
        imap_right.setSpacing(4)
        imap_right.addWidget(FieldLabel("端口"))
        self.imap_port = QLineEdit()
        self.imap_port.setPlaceholderText("993")
        self.imap_port.setText("993")
        self.imap_port.setFixedWidth(100)
        imap_right.addWidget(self.imap_port)
        imap_row1.addLayout(imap_right)
        imap_layout.addLayout(imap_row1)

        imap_row2 = QHBoxLayout()
        imap_row2.setSpacing(10)
        imap_acc = QVBoxLayout()
        imap_acc.setSpacing(4)
        imap_acc.addWidget(FieldLabel("邮箱账号"))
        self.imap_account = QLineEdit()
        self.imap_account.setPlaceholderText("your@email.com")
        imap_acc.addWidget(self.imap_account)
        imap_row2.addLayout(imap_acc)
        imap_pwd = QVBoxLayout()
        imap_pwd.setSpacing(4)
        imap_pwd.addWidget(FieldLabel("密码/授权码"))
        self.imap_password = QLineEdit()
        self.imap_password.setEchoMode(QLineEdit.EchoMode.Password)
        self.imap_password.setPlaceholderText("请输入密码或应用授权码")
        imap_pwd.addWidget(self.imap_password)
        imap_row2.addLayout(imap_pwd)
        imap_layout.addLayout(imap_row2)

        imap_actions = QHBoxLayout()
        self.show_imap_password = QCheckBox("显示授权码")
        self.show_imap_password.toggled.connect(
            lambda checked: self.imap_password.setEchoMode(
                QLineEdit.EchoMode.Normal if checked else QLineEdit.EchoMode.Password
            )
        )
        imap_actions.addWidget(self.show_imap_password)
        imap_actions.addStretch()
        self.imap_test_btn = WhiteButton("🔌 测试连接")
        self.imap_test_btn.setMinimumHeight(42)
        self.imap_test_btn.clicked.connect(self._test_imap_connection)
        imap_actions.addWidget(self.imap_test_btn)
        imap_layout.addLayout(imap_actions)

        scl.addWidget(self.imap_panel)

        layout.addWidget(source_card)

        steps = QHBoxLayout()
        steps.setSpacing(8)
        steps.addWidget(StepBadge("1 选邮件", C_PRIMARY))
        steps.addWidget(StepBadge("2 选文件", C_ACCENT_BLUE))
        steps.addWidget(StepBadge("3 自动收集", C_ACCENT_ORANGE))
        steps.addStretch()
        layout.addLayout(steps)

        card1 = Card()
        cl1 = QVBoxLayout(card1)
        cl1.setContentsMargins(24, 22, 24, 22)
        cl1.setSpacing(14)

        cl1.addWidget(SectionLabel("邮件筛选条件"))

        cl1.addWidget(FieldLabel("邮件标题"))
        title_mode_row = QHBoxLayout()
        title_mode_row.setSpacing(8)

        self.title_mode = SegmentedControl([
            ("contains", "包含文字"),
            ("startswith", "以文字开头"),
        ], default_value="contains")

        title_mode_row.addWidget(self.title_mode)

        self.title_text = QLineEdit()
        self.title_text.setPlaceholderText("例如：Invoice、账单、报价单")
        title_mode_row.addWidget(self.title_text, 1)

        cl1.addLayout(title_mode_row)

        cl1.addWidget(FieldLabel("邮件正文关键词"))
        body_row = QHBoxLayout()
        body_row.setSpacing(8)

        self.body_mode = SegmentedControl([
            ("any", "任一关键词"),
            ("all", "全部关键词"),
        ], default_value="any")

        body_row.addWidget(self.body_mode)

        self.body_text = QLineEdit()
        self.body_text.setPlaceholderText("多个关键词用英文逗号分隔，例如：invoice, payment, 账单")
        body_row.addWidget(self.body_text, 1)

        cl1.addLayout(body_row)

        row2 = QHBoxLayout()
        row2.setSpacing(12)

        left = QVBoxLayout()
        left.setSpacing(6)
        left.addWidget(FieldLabel("发件人包含"))
        self.sender_text = QLineEdit()
        self.sender_text.setPlaceholderText("例如：@supplier.example")
        left.addWidget(self.sender_text)
        row2.addLayout(left)

        right = QVBoxLayout()
        right.setSpacing(6)
        right.addWidget(FieldLabel("收件人包含"))
        self.recipient_text = QLineEdit()
        self.recipient_text.setPlaceholderText("例如：@company.example")
        right.addWidget(self.recipient_text)
        row2.addLayout(right)

        cl1.addLayout(row2)

        cl1.addWidget(FieldLabel("邮件日期范围"))
        self.date_picker = DateRangePicker()
        cl1.addWidget(self.date_picker)

        self.has_attachment = QCheckBox("仅筛选有附件的邮件")
        self.has_attachment.setChecked(True)
        self.has_attachment.setStyleSheet(f"""
            QCheckBox {{
                font-size: 14px;
                font-weight: 700;
                color: {C_TEXT};
                padding-top: 4px;
            }}
        """)
        cl1.addWidget(self.has_attachment)

        layout.addWidget(card1)

        card2 = Card()
        cl2 = QVBoxLayout(card2)
        cl2.setContentsMargins(24, 22, 24, 22)
        cl2.setSpacing(14)

        cl2.addWidget(SectionLabel("收集内容与保存方式"))

        cl2.addWidget(FieldLabel("你想收集什么"))
        self.action_group = QButtonGroup(self)
        act_row = QHBoxLayout()
        act_row.setSpacing(10)

        self.action_attachment = QRadioButton("下载附件")
        self.action_email = QRadioButton("导出整封邮件（.msg / .eml）")
        self.action_both = QRadioButton("附件和邮件都要")
        self.action_both.setChecked(True)

        for rb in [self.action_attachment, self.action_email, self.action_both]:
            rb.setStyleSheet(f"""
                QRadioButton {{
                    font-size: 14px;
                    font-weight: 700;
                    color: {C_TEXT};
                    background: white;
                    border: 2px solid {C_BORDER};
                    border-radius: 16px;
                    padding: 10px 14px;
                }}
                QRadioButton:hover {{
                    border-color: {C_PRIMARY};
                    background: #F8FFF3;
                }}
            """)
            self.action_group.addButton(rb)
            act_row.addWidget(rb)

        act_row.addStretch()
        cl2.addLayout(act_row)

        cl2.addWidget(FieldLabel("附件类型"))
        ext_row = QHBoxLayout()
        ext_row.setSpacing(8)

        self.attachment_ext_presets = QCheckBox(".pdf")
        self.attachment_ext_presets.setChecked(True)
        ext_row.addWidget(self.attachment_ext_presets)

        self.attachment_ext_doc = QCheckBox(".doc/.docx")
        ext_row.addWidget(self.attachment_ext_doc)

        self.attachment_ext_xls = QCheckBox(".xls/.xlsx")
        ext_row.addWidget(self.attachment_ext_xls)

        self.attachment_ext_zip = QCheckBox(".zip/.rar/.7z")
        self.attachment_ext_zip.setChecked(True)
        ext_row.addWidget(self.attachment_ext_zip)

        self.attachment_ext_img = QCheckBox(".png/.jpg")
        ext_row.addWidget(self.attachment_ext_img)

        for cb in [
            self.attachment_ext_presets,
            self.attachment_ext_doc,
            self.attachment_ext_xls,
            self.attachment_ext_zip,
            self.attachment_ext_img
        ]:
            self._style_chip(cb)

        ext_row.addStretch()
        cl2.addLayout(ext_row)

        custom_row = QHBoxLayout()
        custom_row.setSpacing(8)
        custom_row.addWidget(FieldLabel("自定义后缀"))

        self.attachment_ext_custom = QLineEdit()
        self.attachment_ext_custom.setPlaceholderText(".msg, .txt，英文逗号分隔；留空表示不限")
        self.attachment_ext_custom.setFixedWidth(360)
        custom_row.addWidget(self.attachment_ext_custom)
        custom_row.addStretch()

        cl2.addLayout(custom_row)

        attname_row = QHBoxLayout()
        attname_row.setSpacing(8)
        attname_row.addWidget(FieldLabel("附件名称包含"))
        self.attachment_name_text = QLineEdit()
        self.attachment_name_text.setPlaceholderText("合同, 报价单，逗号分隔关键词；留空表示不限")
        self.attachment_name_text.setFixedWidth(360)
        attname_row.addWidget(self.attachment_name_text)
        attname_row.addStretch()
        cl2.addLayout(attname_row)

        unzip_row = QHBoxLayout()
        self.auto_unzip = QCheckBox("自动解压压缩包 .zip / .rar / .7z")
        self.auto_unzip.setChecked(True)
        self.auto_unzip.setStyleSheet(f"""
            QCheckBox {{
                font-size: 14px;
                font-weight: 700;
                color: {C_TEXT};
            }}
        """)
        unzip_row.addWidget(self.auto_unzip)
        unzip_row.addStretch()
        cl2.addLayout(unzip_row)

        cl2.addWidget(FieldLabel("保存位置"))
        out_row = QHBoxLayout()
        out_row.setSpacing(8)

        self.output_dir_edit = QLineEdit()
        self.output_dir_edit.setText(str(DEFAULT_OUTPUT))
        out_row.addWidget(self.output_dir_edit)

        browse_btn = QPushButton("📁 选择")
        browse_btn.setMinimumHeight(42)
        browse_btn.setStyleSheet(f"""
            QPushButton {{
                background: white;
                color: {C_TEXT};
                border: 2px solid {C_BORDER};
                border-radius: 14px;
                padding: 9px 18px;
                font-size: 13px;
                font-weight: 800;
            }}
            QPushButton:hover {{
                border-color: {C_PRIMARY};
                color: {C_PRIMARY};
                background: #F8FFF3;
            }}
        """)
        browse_btn.clicked.connect(self._browse_output)
        out_row.addWidget(browse_btn)

        cl2.addLayout(out_row)

        layout.addWidget(card2)

        btn_row = QHBoxLayout()
        btn_row.addStretch()

        save_plan_btn = WhiteButton("💾 保存为方案")
        save_plan_btn.clicked.connect(self._save_plan)
        btn_row.addWidget(save_plan_btn)

        run_btn = GreenButton("🚀 立即执行")
        run_btn.clicked.connect(self._run)
        btn_row.addWidget(run_btn)

        layout.addLayout(btn_row)

        self.progress = QProgressBar()
        self.progress.setVisible(False)
        self.progress.setFixedHeight(10)
        layout.addWidget(self.progress)

        self.status_label = QLabel("")
        self.status_label.setStyleSheet(f"color: {C_TEXT_SECONDARY}; font-size: 12px;")
        self.status_label.setVisible(False)
        layout.addWidget(self.status_label)

        layout.addStretch()

        scroll.setWidget(container)

        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.addWidget(scroll)

    def _style_chip(self, cb):
        cb.setStyleSheet(f"""
            QCheckBox {{
                background: white;
                color: {C_TEXT_SECONDARY};
                border: 2px solid {C_BORDER};
                border-radius: 18px;
                padding: 7px 12px;
                font-size: 13px;
                font-weight: 700;
            }}
            QCheckBox:hover {{
                border-color: {C_PRIMARY};
                color: {C_PRIMARY};
                background: #F8FFF3;
            }}
            QCheckBox:checked {{
                background: {C_PRIMARY};
                color: white;
                border-color: {C_PRIMARY};
            }}
            QCheckBox::indicator {{
                width: 0px;
                height: 0px;
            }}
        """)

    def _browse_output(self):
        d = QFileDialog.getExistingDirectory(self, translate_text("选择保存目录"), self.output_dir_edit.text())
        if d:
            self.output_dir_edit.setText(d)

    def _get_extensions(self):
        exts = set()

        if self.attachment_ext_presets.isChecked():
            exts.add(".pdf")

        if self.attachment_ext_doc.isChecked():
            exts.update([".doc", ".docx"])

        if self.attachment_ext_xls.isChecked():
            exts.update([".xls", ".xlsx", ".csv"])

        if self.attachment_ext_zip.isChecked():
            exts.update([".zip", ".rar", ".7z"])

        if self.attachment_ext_img.isChecked():
            exts.update([".png", ".jpg", ".jpeg", ".gif", ".bmp"])

        custom = self.attachment_ext_custom.text().strip()
        if custom:
            for e in custom.split(","):
                e = e.strip().lower()
                if e and not e.startswith("."):
                    e = "." + e
                if e:
                    exts.add(e)

        return list(exts) if exts else []

    def _get_conditions(self):
        date_from, date_to = self.date_picker.get_range()

        return {
            "title_mode": self.title_mode.value(),
            "title_text": self.title_text.text(),
            "body_mode": self.body_mode.value(),
            "body_text": self.body_text.text(),
            "sender_text": self.sender_text.text(),
            "recipient_text": self.recipient_text.text(),
            "date_mode": self.date_picker.value(),
            "date_from": date_from,
            "date_to": date_to,
            "has_attachments": self.has_attachment.isChecked(),
            "attachment_exts": self._get_extensions(),
            "attachment_name_text": self.attachment_name_text.text().strip(),
            "auto_unzip": self.auto_unzip.isChecked(),
        }

    def _get_action(self):
        if self.action_attachment.isChecked():
            return "attachment"
        if self.action_email.isChecked():
            return "email"
        return "both"

    def _on_source_changed(self, value):
        self.imap_panel.setVisible(value in ("qq", "imap"))
        self.imap_provider_row.setVisible(value == "imap")
        subtitle = self.findChild(QLabel, "subtitle")
        if value == "outlook":
            subtitle.setText(translate_text("从 Outlook 自动收集邮件、附件和压缩包"))
        elif value == "qq":
            subtitle.setText(translate_text("通过官方 IMAP 安全收集 QQ 邮箱邮件与附件"))
            self._apply_imap_preset("qq")
        else:
            subtitle.setText(translate_text("从支持 IMAP 的邮箱自动收集邮件与附件"))
            self._on_provider_changed()

    def _on_provider_changed(self, *_):
        if self.mail_source.value() != "imap":
            return
        self._apply_imap_preset(self.imap_provider.currentData() or "custom")

    def _apply_imap_preset(self, provider):
        preset = IMAP_PRESETS.get(provider, IMAP_PRESETS["custom"])
        self.imap_server.setText(preset["server"])
        self.imap_port.setText(str(preset["port"]))
        self.imap_server.setReadOnly(provider != "custom")
        self.imap_port.setReadOnly(provider != "custom")
        self.imap_account.setPlaceholderText(translate_text(preset["account_hint"]))
        password_hint = preset.get("password_hint_en") if CURRENT_LANGUAGE == "en_US" else preset["password_hint"]
        help_text = preset.get("help_en") if CURRENT_LANGUAGE == "en_US" else preset["help"]
        security_note = " App passwords are stored only in Windows Credential Manager." if CURRENT_LANGUAGE == "en_US" else "  授权码只保存到 Windows 凭据库。"
        self.imap_password.setPlaceholderText(password_hint)
        self.imap_help.setText("🛡  " + help_text + security_note)

    def _get_imap_config(self, include_password=True):
        if self.mail_source.value() not in ("qq", "imap"):
            return None
        port_text = self.imap_port.text().strip() or "993"
        if not port_text.isdigit() or not (1 <= int(port_text) <= 65535):
            raise ValueError("IMAP 端口必须是 1-65535 之间的数字")
        account = self.imap_account.text().strip()
        password = self.imap_password.text() or load_imap_secret(account)
        return {
            "provider": "qq" if self.mail_source.value() == "qq" else (self.imap_provider.currentData() or "custom"),
            "server": self.imap_server.text().strip(),
            "port": int(port_text),
            "account": account,
            **({"password": password} if include_password else {}),
        }

    def _test_imap_connection(self):
        try:
            config = self._get_imap_config()
        except ValueError as exc:
            QMessageBox.warning(self, "配置有误", str(exc))
            return
        if not config or not config["server"] or not config["account"] or not config.get("password"):
            QMessageBox.warning(self, "信息不完整", "请先填写邮箱账号和授权码。")
            return

        self.imap_test_btn.setEnabled(False)
        self.imap_test_btn.setText(translate_text("正在连接..."))
        self.connection_worker = IMAPConnectionWorker(config)
        self.connection_worker.succeeded.connect(self._on_imap_test_succeeded)
        self.connection_worker.error.connect(self._on_imap_test_failed)
        self.connection_worker.start()

    def _on_imap_test_succeeded(self, message):
        self.imap_test_btn.setEnabled(True)
        self.imap_test_btn.setText(translate_text("🔌 测试连接"))
        QMessageBox.information(self, "IMAP 连接正常", message)

    def _on_imap_test_failed(self, message):
        self.imap_test_btn.setEnabled(True)
        self.imap_test_btn.setText(translate_text("🔌 测试连接"))
        QMessageBox.critical(self, "IMAP 连接失败", message)

    def _run(self):
        output = self.output_dir_edit.text().strip()

        if not output:
            QMessageBox.warning(self, "提示", "请选择保存目录")
            return

        os.makedirs(output, exist_ok=True)

        conditions = refresh_dynamic_dates(self._get_conditions())

        if self.mail_source.value() in ("qq", "imap"):
            try:
                imap_config = self._get_imap_config()
            except ValueError as exc:
                QMessageBox.warning(self, "配置有误", str(exc))
                return
            if not imap_config or not imap_config["server"] or not imap_config["account"] or not imap_config.get("password"):
                QMessageBox.warning(self, "提示", "请填写完整的 IMAP 配置和授权码")
                return
            self.worker = IMAPWorker(conditions, output, self._get_action(), imap_config)
        else:
            self.worker = OutlookWorker(conditions, output, self._get_action())

        self.progress.setVisible(True)
        self.progress.setRange(0, 0)
        self.status_label.setVisible(True)
        self.status_label.setText("正在初始化...")

        self.worker.progress.connect(self._on_progress)
        self.worker.finished.connect(self._on_finished)
        self.worker.error.connect(self._on_error)
        self.worker.start()

    def _on_progress(self, msg):
        self.status_label.setText(translate_text(msg))

    def _on_finished(self, result):
        self.progress.setRange(0, 100)
        self.progress.setValue(100)
        self.status_label.setText("✅ 完成!")

        conditions = result.get("conditions", {})
        fail_reasons = result.get("fail_reasons", {})

        msg = "━━━ 搜索条件 ━━━\n"
        msg += f"标题: {conditions.get('title_text') or '(未设置)'}  [{conditions.get('title_mode', 'contains')}]\n"
        msg += f"正文: {conditions.get('body_text') or '(未设置)'}  [{conditions.get('body_mode', 'any')}]\n"
        msg += f"发件人: {conditions.get('sender_text') or '(未设置)'}\n"
        msg += f"收件人: {conditions.get('recipient_text') or '(未设置)'}\n"

        if conditions.get("date_from"):
            msg += f"日期从: {conditions['date_from']}\n"
        else:
            msg += "日期从: 不限\n"

        if conditions.get("date_to"):
            msg += f"日期到: {conditions['date_to']}\n"
        else:
            msg += "日期到: 不限\n"

        msg += f"仅附件: {'是' if conditions.get('has_attachments') else '否'}\n"
        msg += f"附件类型: {', '.join(conditions.get('attachment_exts', [])) or '(不限)'}\n"
        msg += f"自动解压: {'是' if conditions.get('auto_unzip') else '否'}\n"

        msg += "\n━━━ 扫描结果 ━━━\n"
        msg += f"扫描邮件: {result['scanned']} 封\n"
        msg += f"匹配邮件: {result['matched']} 封\n"
        msg += f"保存文件: {result['saved']} 个\n"
        msg += f"跳过重复: {result.get('skipped_duplicate', 0)} 个\n"
        msg += f"保存位置: {result['output']}\n"

        if fail_reasons:
            msg += "\n━━━ 不匹配原因统计 ━━━\n"
            for r, cnt in sorted(fail_reasons.items(), key=lambda x: -x[1]):
                msg += f"  {r}: {cnt} 封\n"

        if result["errors"]:
            msg += f"\n⚠ 错误: {len(result['errors'])} 个\n" + "\n".join(result["errors"][:3])

        QMessageBox.information(self, "任务完成", msg)

        self.progress.setVisible(False)
        self.status_label.setVisible(False)

    def _on_error(self, err):
        self.progress.setVisible(False)
        self.status_label.setVisible(False)
        QMessageBox.critical(self, "错误", f"执行失败:\n{err}")

    def _save_plan(self):
        editing = bool(self.editing_plan_id)

        dlg = SavePlanDialog(
            default_name=self.editing_plan_name if editing else "",
            editing=editing,
            parent=self
        )

        if dlg.exec() != QDialog.DialogCode.Accepted:
            return

        mode, name = dlg.result_data()

        if not name:
            QMessageBox.warning(self, "提示", "请输入方案名称")
            return

        plans = load_plans()

        old_schedule = {
            "enabled": False,
            "mode": "daily",
            "time": "09:00",
            "months": 1,
            "end_date": "",
            "archive_mode": "by_date",
            "last_run": "",
            "last_archive_date": "",
        }
        old_created = datetime.now().strftime("%Y-%m-%d %H:%M")

        if self.editing_plan_id:
            for p in plans:
                if p.get("id") == self.editing_plan_id:
                    old_schedule = p.get("schedule", old_schedule)
                    old_created = p.get("created", old_created)
                    break

        plan_id = self.editing_plan_id if mode == "overwrite" and self.editing_plan_id else new_plan_id()
        imap_config = None
        if self.mail_source.value() in ("qq", "imap"):
            try:
                runtime_config = self._get_imap_config(include_password=True)
            except ValueError as exc:
                QMessageBox.warning(self, "配置有误", str(exc))
                return
            if not runtime_config or not runtime_config.get("account"):
                QMessageBox.warning(self, "信息不完整", "请先填写邮箱账号。")
                return
            secret = runtime_config.pop("password", "")
            if secret and not save_imap_secret(runtime_config["account"], secret):
                QMessageBox.warning(
                    self,
                    "授权码未保存",
                    "无法写入系统凭据库。方案仍会保存，但下次运行需重新输入授权码。\n"
                    "请重新运行「安装依赖.bat」后再试。",
                )
            imap_config = runtime_config

        data = {
            "id": plan_id,
            "name": name,
            "conditions": self._get_conditions(),
            "action": self._get_action(),
            "output_dir": self.output_dir_edit.text(),
            "mail_source": self.mail_source.value(),
            "imap_config": imap_config,
            "created": old_created if mode == "overwrite" and self.editing_plan_id else datetime.now().strftime("%Y-%m-%d %H:%M"),
            "updated": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "schedule": old_schedule if mode == "overwrite" and self.editing_plan_id else {
                "enabled": False,
                "mode": "daily",
                "time": "09:00",
                "months": 1,
                "end_date": "",
                "archive_mode": "by_date",
                "last_run": "",
                "last_archive_date": "",
            }
        }

        if mode == "overwrite" and self.editing_plan_id:
            plans = [data if p.get("id") == self.editing_plan_id else p for p in plans]
            QMessageBox.information(self, "已覆盖", f"方案「{name}」已覆盖保存。")
        else:
            plans.append(data)
            QMessageBox.information(self, "已保存", f"方案「{name}」已另存为新方案。")

        save_plans(plans)

        self.editing_plan_id = data["id"]
        self.editing_plan_name = name

        self.main.plans_page.refresh()

    def load_plan_for_edit(self, plan):
        self.editing_plan_id = plan.get("id")
        self.editing_plan_name = plan.get("name", "")

        c = plan.get("conditions", {})

        self.title_mode.set_value(
            "startswith" if c.get("title_mode") == "startswith" else "contains"
        )
        self.title_text.setText(c.get("title_text", ""))

        self.body_mode.set_value(
            "all" if c.get("body_mode") == "all" else "any"
        )
        self.body_text.setText(c.get("body_text", ""))

        self.sender_text.setText(c.get("sender_text", ""))
        self.recipient_text.setText(c.get("recipient_text", ""))

        date_mode = c.get("date_mode")
        if date_mode and date_mode != "custom":
            self.date_picker.set_value(date_mode)
        else:
            self.date_picker.set_range(c.get("date_from"), c.get("date_to"))

        self.has_attachment.setChecked(c.get("has_attachments", True))
        self.output_dir_edit.setText(plan.get("output_dir", str(DEFAULT_OUTPUT)))

        exts = c.get("attachment_exts", [])
        exts_lower = [e.lower() for e in exts]

        self.attachment_ext_presets.setChecked(".pdf" in exts_lower)
        self.attachment_ext_doc.setChecked(any(e in exts_lower for e in [".doc", ".docx"]))
        self.attachment_ext_xls.setChecked(any(e in exts_lower for e in [".xls", ".xlsx", ".csv"]))
        self.attachment_ext_zip.setChecked(any(e in exts_lower for e in [".zip", ".rar", ".7z"]))
        self.attachment_ext_img.setChecked(any(e in exts_lower for e in [".png", ".jpg", ".jpeg", ".gif", ".bmp"]))

        preset_all = {
            ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".csv",
            ".zip", ".rar", ".7z", ".png", ".jpg", ".jpeg",
            ".gif", ".bmp"
        }

        custom = [e for e in exts if e.lower() not in preset_all]
        self.attachment_ext_custom.setText(", ".join(custom))

        self.attachment_name_text.setText(c.get("attachment_name_text", ""))
        self.auto_unzip.setChecked(c.get("auto_unzip", True))

        act = plan.get("action", "both")
        if act == "attachment":
            self.action_attachment.setChecked(True)
        elif act == "email":
            self.action_email.setChecked(True)
        else:
            self.action_both.setChecked(True)

        # 还原邮箱来源和IMAP配置
        source = plan.get("mail_source", "outlook")
        imap_cfg = plan.get("imap_config") or {}
        if source == "imap" and imap_cfg.get("provider") == "qq":
            source = "qq"
        self.mail_source.set_value(source)
        self.imap_panel.setVisible(source in ("qq", "imap"))
        self.imap_provider_row.setVisible(source == "imap")
        if source in ("qq", "imap"):
            provider = "qq" if source == "qq" else imap_cfg.get("provider", "custom")
            if source == "imap":
                index = self.imap_provider.findData(provider)
                self.imap_provider.setCurrentIndex(index if index >= 0 else self.imap_provider.findData("custom"))
            self._apply_imap_preset(provider)
            self.imap_server.setText(imap_cfg.get("server", ""))
            self.imap_port.setText(str(imap_cfg.get("port", "993")))
            self.imap_account.setText(imap_cfg.get("account", ""))
            self.imap_password.setText(load_imap_secret(imap_cfg.get("account", "")))
            subtitle_text = (
                "通过官方 IMAP 安全收集 QQ 邮箱邮件与附件"
                if source == "qq" else "从支持 IMAP 的邮箱自动收集邮件与附件"
            )
            self.findChild(QLabel, "subtitle").setText(translate_text(subtitle_text))

    def clear_editing_plan(self):
        self.editing_plan_id = None
        self.editing_plan_name = ""


# ─── 批量附件下载工作线程 ───

class BatchDownloadWorker(QThread):
    progress = Signal(str)
    finished = Signal(dict)
    error = Signal(str)

    def __init__(self, target_names, output_dir, mail_source, imap_config=None, exact_match=False):
        super().__init__()
        self.target_names = [n.strip().lower() for n in target_names if n.strip()]
        self.output_dir = output_dir
        self.mail_source = mail_source
        self.imap_config = imap_config
        self.exact_match = exact_match

    def run(self):
        try:
            if self.mail_source == "outlook":
                self._run_outlook()
            else:
                self._run_imap()
        except Exception as e:
            self.error.emit(str(e))

    def _match_name(self, fname):
        fn = fname.lower()
        if self.exact_match:
            return fn in self.target_names
        return any(t in fn for t in self.target_names)

    def _run_outlook(self):
        if win32com is None:
            raise RuntimeError("未检测到 pywin32。请先安装：pip install pywin32")

        outlook = win32com.client.Dispatch("Outlook.Application")
        namespace = outlook.GetNamespace("MAPI")
        inbox = namespace.GetDefaultFolder(6)

        self.progress.emit("🔍 正在搜索邮件...")
        messages = inbox.Items
        messages.Sort("[ReceivedTime]", True)
        total = messages.Count
        self.progress.emit(f"📬 扫描 {total} 封邮件...")

        saved = 0
        errors = []

        for i in range(1, total + 1):
            try:
                mail = messages.Item(i)
            except Exception:
                continue

            try:
                atts = mail.Attachments
            except Exception:
                continue

            for att in atts:
                fname = att.FileName
                if not self._match_name(fname):
                    continue
                save_path = os.path.join(self.output_dir, fname)
                base, e = os.path.splitext(fname)
                c = 1
                while os.path.exists(save_path):
                    save_path = os.path.join(self.output_dir, f"{base}_{c}{e}")
                    c += 1
                try:
                    att.SaveAsFile(save_path)
                    saved += 1
                    self.progress.emit(f"  📎 保存: {os.path.basename(save_path)}")
                except Exception as ex:
                    errors.append(f"{fname}: {ex}")

        result = {"saved": saved, "errors": errors, "output": self.output_dir}
        self.finished.emit(result)

    def _run_imap(self):
        cfg = self.imap_config or {}
        server = cfg.get("server", "")
        port = cfg.get("port", 993)
        account = cfg.get("account", "")
        password = cfg.get("password", "")

        if not server or not account or not password:
            raise RuntimeError("IMAP 配置不完整")

        conn = imaplib.IMAP4_SSL(server, int(port))
        conn.login(account, password)
        conn.select("INBOX")

        self.progress.emit("🔍 正在搜索邮件...")
        status, data = conn.search(None, "ALL")
        if status != "OK":
            raise RuntimeError("服务器搜索邮件失败")
        mail_ids = list(reversed(data[0].split()))
        total = len(mail_ids)
        self.progress.emit(f"📬 待检查邮件: {total} 封")

        saved = 0
        errors = []

        for mid in mail_ids:
            try:
                fetch_status, msg_data = conn.fetch(mid, "(RFC822)")
                if fetch_status != "OK" or not msg_data or not isinstance(msg_data[0], tuple):
                    continue
                raw_email = msg_data[0][1]
                msg = email.message_from_bytes(raw_email)
            except Exception:
                continue

            for part in msg.walk():
                fname = part.get_filename()
                if not fname:
                    continue
                fname = self._sanitize_filename(self._decode_header(fname))
                if not self._match_name(fname):
                    continue
                save_path = os.path.join(self.output_dir, fname)
                base, e = os.path.splitext(fname)
                c = 1
                while os.path.exists(save_path):
                    save_path = os.path.join(self.output_dir, f"{base}_{c}{e}")
                    c += 1
                try:
                    with open(save_path, "wb") as f:
                        f.write(part.get_payload(decode=True))
                    saved += 1
                    self.progress.emit(f"  📎 保存: {os.path.basename(save_path)}")
                except Exception as ex:
                    errors.append(f"{fname}: {ex}")

        conn.logout()
        result = {"saved": saved, "errors": errors, "output": self.output_dir}
        self.finished.emit(result)

    @staticmethod
    def _sanitize_filename(s):
        return re.sub(r'[\\/*?:"<>|]', "_", s)[:100]

    @staticmethod
    def _decode_header(s):
        if not s:
            return ""
        try:
            from email.header import decode_header
            parts = decode_header(s) if isinstance(s, str) else []
            result = ""
            for part, charset in parts:
                if isinstance(part, bytes):
                    result += part.decode(charset or "utf-8", errors="replace")
                else:
                    result += part
            return result
        except Exception:
            return str(s)


# ─── 批量下载页面 ───

class BatchDownloadPage(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main = main_window
        self.setMinimumWidth(600)

        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet(f"QScrollArea {{ background: {C_CARD}; border: none; }}")

        container = QWidget()
        layout = QVBoxLayout(container)
        layout.setSpacing(16)
        layout.setContentsMargins(40, 30, 40, 30)

        header = QHBoxLayout()
        icon_lbl = QLabel("📦")
        icon_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon_lbl.setFixedSize(50, 50)
        icon_lbl.setStyleSheet("background: #FFF4D6; border-radius: 25px; font-size: 24px;")
        header.addWidget(icon_lbl)
        title = SectionLabel("批量附件下载")
        header.addWidget(title)
        header.addStretch()
        layout.addLayout(header)

        source_card = Card()
        sc = QVBoxLayout(source_card)
        sc.setContentsMargins(18, 14, 18, 14)
        sc.setSpacing(12)

        source_header = QHBoxLayout()
        source_header.addWidget(FieldLabel("邮箱来源"))
        self.mail_source = SegmentedControl([
            ("outlook", "Outlook"),
            ("qq", "QQ 邮箱"),
            ("imap", "IMAP"),
        ], "outlook")
        self.mail_source.value_changed.connect(self._on_source_changed)
        source_header.addWidget(self.mail_source)
        source_header.addStretch()
        sc.addLayout(source_header)

        self.imap_stack = QStackedWidget()
        self.imap_stack.setStyleSheet("background: transparent;")
        self.no_imap = QWidget()
        self.imap_stack.addWidget(self.no_imap)

        self.imap_config_widget = QWidget()
        icw = QVBoxLayout(self.imap_config_widget)
        icw.setContentsMargins(0, 8, 0, 0)
        icw.setSpacing(10)

        prov_row = QHBoxLayout()
        prov_row.addWidget(FieldLabel("服务商"))
        self.imap_provider = QComboBox()
        self.imap_provider.setMinimumWidth(200)
        self.imap_provider.addItem("QQ 邮箱", "qq")
        self.imap_provider.addItem("163 邮箱", "163")
        self.imap_provider.addItem("Gmail", "gmail")
        self.imap_provider.addItem("其他", "custom")
        self.imap_provider.currentIndexChanged.connect(self._on_provider_changed)
        prov_row.addWidget(self.imap_provider)
        prov_row.addStretch()
        icw.addLayout(prov_row)

        for label, attr_name, ph, width in [
            ("邮箱账号", "imap_account", "your@email.com", 360),
            ("授权码", "imap_password", "邮箱授权码（非邮箱密码）", 360),
            ("服务器", "imap_server", "imap.qq.com", 360),
        ]:
            r = QHBoxLayout()
            r.addWidget(FieldLabel(label))
            le = QLineEdit()
            le.setPlaceholderText(ph)
            le.setFixedWidth(width)
            if attr_name == "imap_password":
                le.setEchoMode(QLineEdit.EchoMode.Password)
            setattr(self, attr_name, le)
            r.addWidget(le)
            r.addStretch()
            icw.addLayout(r)

        port_row = QHBoxLayout()
        port_row.addWidget(FieldLabel("端口"))
        self.imap_port = QSpinBox()
        self.imap_port.setMinimum(1)
        self.imap_port.setMaximum(65535)
        self.imap_port.setValue(993)
        self.imap_port.setFixedWidth(100)
        port_row.addWidget(self.imap_port)
        port_row.addStretch()
        icw.addLayout(port_row)

        self.imap_stack.addWidget(self.imap_config_widget)
        sc.addWidget(self.imap_stack)
        layout.addWidget(source_card)

        name_card = Card()
        nc = QVBoxLayout(name_card)
        nc.setContentsMargins(18, 14, 18, 14)
        nc.setSpacing(12)
        nc.addWidget(FieldLabel("要下载的附件名（每行一个，或逗号分隔）"))
        self.name_edit = QTextEdit()
        self.name_edit.setPlaceholderText("合同.pdf\n报价单.xlsx\n附件001.docx\n...")
        self.name_edit.setMinimumHeight(120)
        nc.addWidget(self.name_edit)

        btns_row = QHBoxLayout()
        btns_row.setSpacing(8)
        self.import_csv_btn = WhiteButton("📄 导入 Excel/CSV")
        self.import_csv_btn.clicked.connect(self._import_file)
        btns_row.addWidget(self.import_csv_btn)
        self.download_template_btn = WhiteButton("📥 下载模板")
        self.download_template_btn.clicked.connect(self._download_template)
        btns_row.addWidget(self.download_template_btn)
        self.exact_cb = QCheckBox("精确匹配（文件名完全相等，否则为模糊匹配）")
        self.exact_cb.setStyleSheet(f"color: {C_TEXT_SECONDARY}; font-size: 13px;")
        btns_row.addWidget(self.exact_cb)
        btns_row.addStretch()
        nc.addLayout(btns_row)
        layout.addWidget(name_card)

        action_card = Card()
        ac = QVBoxLayout(action_card)
        ac.setContentsMargins(18, 14, 18, 14)
        ac.setSpacing(10)
        out_row = QHBoxLayout()
        out_row.addWidget(FieldLabel("输出目录"))
        self.output_dir_edit = QLineEdit()
        self.output_dir_edit.setText(str(DEFAULT_OUTPUT))
        self.output_dir_edit.setMinimumWidth(300)
        out_row.addWidget(self.output_dir_edit)
        self.browse_btn = WhiteButton("浏览")
        self.browse_btn.clicked.connect(self._browse_output)
        out_row.addWidget(self.browse_btn)
        out_row.addStretch()
        ac.addLayout(out_row)

        run_row = QHBoxLayout()
        self.run_btn = GreenButton("🚀 开始批量下载")
        self.run_btn.clicked.connect(self._run)
        run_row.addWidget(self.run_btn)
        run_row.addStretch()
        ac.addLayout(run_row)

        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        ac.addWidget(self.progress_bar)

        self.status_label = QLabel("")
        self.status_label.setStyleSheet(f"color: {C_TEXT_SECONDARY}; font-size: 13px;")
        self.status_label.setVisible(False)
        ac.addWidget(self.status_label)
        layout.addWidget(action_card)
        layout.addStretch()

        scroll.setWidget(container)
        outer.addWidget(scroll)
        self._on_source_changed("outlook")

    def _on_source_changed(self, value):
        self.imap_stack.setVisible(value != "outlook")
        if value == "outlook":
            self.imap_stack.setCurrentIndex(0)
        else:
            self.imap_stack.setCurrentIndex(1)
            self._on_provider_changed(0)

    def _on_provider_changed(self, idx):
        provider = self.imap_provider.currentData()
        preset = IMAP_PRESETS.get(provider, IMAP_PRESETS["custom"])
        self.imap_server.setText(preset["server"])
        self.imap_port.setValue(preset["port"])
        self.imap_account.setPlaceholderText(preset["account_hint"])
        self.imap_password.setPlaceholderText(preset["password_hint"])

    def _browse_output(self):
        path = QFileDialog.getExistingDirectory(self, "选择输出目录")
        if path:
            self.output_dir_edit.setText(path)

    def _import_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "导入文件", "", "Excel/CSV 文件 (*.xlsx *.xls *.csv);;所有文件 (*)"
        )
        if not path:
            return
        names = []
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == ".csv":
                import csv
                with open(path, "r", encoding="utf-8-sig") as f:
                    reader = csv.reader(f)
                    for row in reader:
                        for cell in row:
                            cell = cell.strip()
                            if cell:
                                names.append(cell)
            else:
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(path, read_only=True)
                    ws = wb.active
                    for row in ws.iter_rows(values_only=True):
                        for cell in row:
                            if cell and str(cell).strip():
                                names.append(str(cell).strip())
                except ImportError:
                    QMessageBox.warning(self, "缺少依赖", "需要 openpyxl 库才能读取 Excel 文件。\n请执行：pip install openpyxl")
                    return
        except Exception as e:
            QMessageBox.warning(self, "导入失败", f"读取文件失败：{e}")
            return
        existing = set()
        for line in self.name_edit.toPlainText().strip().split("\n"):
            line = line.strip()
            if line:
                existing.add(line)
        for n in names:
            if n not in existing:
                existing.add(n)
        self.name_edit.setPlainText("\n".join(sorted(existing)))
        self.status_label.setText(f"✅ 已导入 {len(names)} 个文件名（去重合并）")
        self.status_label.setVisible(True)

    def _download_template(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "保存模板", "批量下载模板.xlsx", "Excel 文件 (*.xlsx)"
        )
        if not path:
            return
        try:
            import openpyxl
        except ImportError:
            QMessageBox.warning(self, "缺少依赖", "需要 openpyxl 库。\n请执行：pip install openpyxl")
            return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "附件名列表"
            ws["A1"] = "附件名（每行一个）"
            ws["A1"].font = openpyxl.styles.Font(bold=True, size=12)
            examples = [
                "合同2024-Q1.pdf",
                "报价单_供应商A.xlsx",
                "发票_INV-001.pdf",
                "报告_202406.docx",
                "项目计划_v2.xlsx",
            ]
            for i, name in enumerate(examples, start=2):
                ws[f"A{i}"] = name
            ws.column_dimensions["A"].width = 40
            wb.save(path)
            self.status_label.setText(f"✅ 模板已保存到：{path}")
            self.status_label.setVisible(True)
        except Exception as e:
            QMessageBox.warning(self, "保存失败", str(e))

    def _run(self):
        raw = self.name_edit.toPlainText().strip()
        lines = []
        for line in raw.split("\n"):
            line = line.strip()
            if line:
                if "," in line and "\n" not in raw and len(list(raw.split("\n"))) <= 1:
                    for item in line.split(","):
                        item = item.strip()
                        if item:
                            lines.append(item)
                else:
                    lines.append(line)
        if not lines:
            QMessageBox.warning(self, "无附件名", "请先输入要下载的附件名称。")
            return
        output = self.output_dir_edit.text().strip()
        if not output:
            QMessageBox.warning(self, "无输出目录", "请先设置输出目录。")
            return
        os.makedirs(output, exist_ok=True)
        mail_source = self.mail_source.value()
        if mail_source == "outlook":
            imap_config = None
        else:
            imap_config = {
                "server": self.imap_server.text().strip(),
                "port": self.imap_port.value(),
                "account": self.imap_account.text().strip(),
                "password": self.imap_password.text(),
                "provider": self.imap_provider.currentData() or "custom",
            }
            if not imap_config["account"] or not imap_config["password"]:
                QMessageBox.warning(self, "IMAP 配置不完整", "请填写邮箱账号和授权码。")
                return
        self.run_btn.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 0)
        self.status_label.setText("⏳ 正在搜索...")
        self.status_label.setVisible(True)

        self.worker = BatchDownloadWorker(
            lines, output, mail_source, imap_config,
            exact_match=self.exact_cb.isChecked()
        )
        self.worker.progress.connect(lambda msg: self.status_label.setText(msg))
        self.worker.finished.connect(self._on_finished)
        self.worker.error.connect(self._on_error)
        self.worker.start()

    def _on_finished(self, result):
        self.run_btn.setEnabled(True)
        self.progress_bar.setVisible(False)
        saved = result.get("saved", 0)
        errors = result.get("errors", [])
        msg = f"✅ 下载完成：{saved} 个附件保存到 {result['output']}"
        if errors:
            msg += f"\n⚠ {len(errors)} 个失败"
        self.status_label.setText(msg)
        QMessageBox.information(self, "批量下载完成", msg)

    def _on_error(self, err):
        self.run_btn.setEnabled(True)
        self.progress_bar.setVisible(False)
        self.status_label.setText(f"❌ 错误：{err}")
        QMessageBox.warning(self, "下载失败", str(err))


# ─── 页面2: 我的方案 ───

class PlansPage(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main = main_window

        layout = QVBoxLayout(self)
        layout.setSpacing(18)
        layout.setContentsMargins(40, 30, 40, 30)

        header = QHBoxLayout()

        icon = QLabel("📚")
        icon.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon.setFixedSize(58, 58)
        icon.setStyleSheet("""
            QLabel {
                background: #E8F5FF;
                border-radius: 29px;
                font-size: 30px;
                border: none;
            }
        """)
        header.addWidget(icon)

        text_box = QVBoxLayout()
        text_box.setSpacing(4)
        text_box.addWidget(TitleLabel("我的方案"))
        text_box.addWidget(SubtitleLabel("保存常用筛选规则，下次一键执行，也可以设置定时运行"))
        header.addLayout(text_box)
        header.addStretch()

        layout.addLayout(header)

        self.plan_list = QListWidget()
        self.plan_list.setSpacing(10)
        self.plan_list.setUniformItemSizes(False)
        self.plan_list.itemClicked.connect(self._on_click)
        layout.addWidget(self.plan_list)

        btn_row = QHBoxLayout()
        btn_row.addStretch()

        run_btn = GreenButton("🚀 执行选中方案")
        run_btn.clicked.connect(self._run_selected)
        btn_row.addWidget(run_btn)

        del_btn = DangerButton("🗑 删除选中")
        del_btn.clicked.connect(self._delete_selected)
        btn_row.addWidget(del_btn)

        layout.addLayout(btn_row)

        self.refresh()

    def refresh(self):
        self.plan_list.clear()
        self.plan_list.setSpacing(10)
        self.plan_list.setUniformItemSizes(False)

        plans = load_plans()

        if not plans:
            item = QListWidgetItem()
            item.setFlags(Qt.ItemFlag.NoItemFlags)
            item.setSizeHint(QSize(0, 96))

            widget = QWidget()
            widget.setStyleSheet(f"""
                QWidget {{
                    background: {C_CARD};
                    border-radius: 18px;
                }}
            """)

            wl = QVBoxLayout(widget)
            wl.setContentsMargins(18, 14, 18, 14)
            wl.setSpacing(6)

            empty_title = QLabel("还没有保存方案")
            empty_title.setMinimumHeight(26)
            empty_title.setFont(QFont("Microsoft YaHei UI", 14, QFont.Bold))
            empty_title.setStyleSheet(f"""
                QLabel {{
                    color: {C_TEXT};
                    background: transparent;
                    border: none;
                }}
            """)
            wl.addWidget(empty_title)

            empty_subtitle = QLabel("在「邮件收集助手」页面设置筛选条件后，可以保存为方案。")
            empty_subtitle.setMinimumHeight(22)
            empty_subtitle.setFont(QFont("Microsoft YaHei UI", 11))
            empty_subtitle.setStyleSheet(f"""
                QLabel {{
                    color: {C_TEXT_SECONDARY};
                    background: transparent;
                    border: none;
                }}
            """)
            wl.addWidget(empty_subtitle)

            self.plan_list.addItem(item)
            self.plan_list.setItemWidget(item, widget)
            retranslate_widget_tree(self)
            return

        for plan in plans:
            item = QListWidgetItem()
            item.setData(Qt.ItemDataRole.UserRole, plan)
            item.setSizeHint(QSize(0, 118))

            widget = QWidget()
            widget.setStyleSheet(f"""
                QWidget {{
                    background: {C_CARD};
                    border-radius: 18px;
                }}
            """)

            outer = QHBoxLayout(widget)
            outer.setContentsMargins(18, 13, 14, 13)
            outer.setSpacing(10)

            # 左侧：名称 + 摘要
            left_box = QVBoxLayout()
            left_box.setSpacing(6)

            name_edit = QLineEdit(plan.get("name", "未命名方案"))
            name_edit.setMinimumHeight(34)
            name_edit.setFont(QFont("Microsoft YaHei UI", 14, QFont.Bold))
            name_edit.setToolTip("点击即可修改方案名称，离开输入框后自动保存")
            name_edit.setStyleSheet(f"""
                QLineEdit {{
                    color: {C_TEXT};
                    background: transparent;
                    border: none;
                    padding: 0px;
                }}
                QLineEdit:focus {{
                    background: white;
                    border: 2px solid {C_PRIMARY};
                    border-radius: 10px;
                    padding: 3px 8px;
                }}
            """)
            name_edit.editingFinished.connect(
                lambda p_id=plan.get("id"), edit=name_edit: self._rename_plan(p_id, edit.text())
            )
            left_box.addWidget(name_edit)

            action_text = action_to_text(plan.get("action"))
            schedule = plan.get("schedule", {})

            if schedule.get("enabled"):
                schedule_text = f"每天 {schedule.get('time', '09:00')} 执行，至 {schedule.get('end_date', '')}"
            else:
                schedule_text = "未定时"

            detail = f"处理方式：{action_text}  ｜  创建：{plan.get('created', '')}  ｜  定时：{schedule_text}"

            detail_lbl = QLabel(detail)
            detail_lbl.setMinimumHeight(24)
            detail_lbl.setFont(QFont("Microsoft YaHei UI", 11))
            detail_lbl.setStyleSheet(f"""
                QLabel {{
                    color: {C_TEXT_SECONDARY};
                    background: transparent;
                    border: none;
                }}
            """)
            left_box.addWidget(detail_lbl)

            outer.addLayout(left_box, 1)

            # 右侧按钮：明细 / 编辑 / 定时
            right_box = QHBoxLayout()
            right_box.setSpacing(6)
            right_box.setAlignment(Qt.AlignmentFlag.AlignCenter)

            detail_btn = SmallActionButton("明细", C_ACCENT_BLUE)
            detail_btn.clicked.connect(lambda checked=False, p=plan: self._show_detail(p))
            right_box.addWidget(detail_btn)

            edit_btn = SmallActionButton("编辑", C_PRIMARY)
            edit_btn.clicked.connect(lambda checked=False, p=plan: self._edit_plan(p))
            right_box.addWidget(edit_btn)

            schedule_btn = QPushButton("⏰ 设置定时")
            schedule_btn.setFixedSize(92, 34)
            schedule_btn.setCursor(Qt.CursorShape.PointingHandCursor)
            schedule_btn.setFont(QFont("Microsoft YaHei UI", 10, QFont.Bold))
            schedule_btn.setStyleSheet(f"""
                QPushButton {{
                    background: white;
                    color: {C_ACCENT_ORANGE};
                    border: 2px solid {C_ACCENT_ORANGE};
                    border-radius: 12px;
                    padding: 0px;
                }}
                QPushButton:hover {{
                    background: #FFF4D6;
                    color: #C87500;
                }}
                QPushButton:pressed {{
                    background: {C_ACCENT_ORANGE};
                    color: white;
                }}
            """)

            schedule_btn.clicked.connect(lambda checked=False, p=plan: self._schedule_plan(p))
            right_box.addWidget(schedule_btn)

            outer.addLayout(right_box)

            self.plan_list.addItem(item)
            self.plan_list.setItemWidget(item, widget)

        retranslate_widget_tree(self)

    def _rename_plan(self, plan_id, new_name):
        new_name = new_name.strip()

        if not new_name:
            return

        plans = load_plans()
        changed = False

        for p in plans:
            if p.get("id") == plan_id and p.get("name") != new_name:
                p["name"] = new_name
                p["updated"] = datetime.now().strftime("%Y-%m-%d %H:%M")
                changed = True
                break

        if changed:
            save_plans(plans)
            self.refresh()

    def _show_detail(self, plan):
        # 重新读取一次，防止名称刚刚被修改但 plan 对象还是旧的
        fresh_plan = plan
        for p in load_plans():
            if p.get("id") == plan.get("id"):
                fresh_plan = p
                break

        dlg = PlanDetailDialog(fresh_plan, self)
        dlg.exec()

    def _edit_plan(self, plan):
        fresh_plan = plan
        for p in load_plans():
            if p.get("id") == plan.get("id"):
                fresh_plan = p
                break

        self.main.create_page.load_plan_for_edit(fresh_plan)
        self.main.nav.setCurrentIndex(0)
        self.main.navbar.set_active(0)

        QMessageBox.information(
            self,
            "已进入编辑",
            "方案已加载到「收集助手」。\n\n修改后点击「保存为方案」，可以选择覆盖原方案或另存为新方案。"
        )

    def _schedule_plan(self, plan):
        fresh_plan = plan
        for p in load_plans():
            if p.get("id") == plan.get("id"):
                fresh_plan = p
                break

        dlg = ScheduleDialog(fresh_plan, self)

        if dlg.exec() != QDialog.DialogCode.Accepted:
            return

        schedule_data = dlg.result_data()

        plans = load_plans()
        for p in plans:
            if p.get("id") == fresh_plan.get("id"):
                p["schedule"] = schedule_data
                p["updated"] = datetime.now().strftime("%Y-%m-%d %H:%M")
                break

        save_plans(plans)
        self.refresh()

        if schedule_data.get("enabled"):
            QMessageBox.information(
                self,
                "定时已启用",
                f"方案「{fresh_plan.get('name')}」将每天 {schedule_data.get('time')} 自动执行。\n\n注意：程序需要保持打开。"
            )
        else:
            QMessageBox.information(
                self,
                "定时已关闭",
                f"方案「{fresh_plan.get('name')}」已关闭定时执行。"
            )

    def _on_click(self, item):
        pass

    def _run_selected(self):
        items = self.plan_list.selectedItems()

        if not items:
            QMessageBox.warning(self, "提示", "请先选择一个方案")
            return

        plan = items[0].data(Qt.ItemDataRole.UserRole)

        if not plan:
            return

        self.main.run_plan(plan, scheduled=False)

    def _delete_selected(self):
        items = self.plan_list.selectedItems()

        if not items:
            QMessageBox.warning(self, "提示", "请先选择一个方案")
            return

        plan = items[0].data(Qt.ItemDataRole.UserRole)

        if not plan:
            return

        reply = QMessageBox.question(
            self,
            "确认删除",
            f"删除方案「{plan.get('name', '')}」？",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )

        if reply == QMessageBox.StandardButton.Yes:
            plans = load_plans()
            plans = [p for p in plans if p.get("id") != plan.get("id")]
            save_plans(plans)
            self.refresh()


# ─── 自定义窗口顶栏与缩放手柄 ───

class WindowTitleBar(QWidget):
    def __init__(self, window):
        super().__init__(window)
        self.window = window
        self.setObjectName("WindowTitleBar")
        self.setFixedHeight(54)

        palette = current_theme()
        border_width = 3 if APP_SETTINGS["theme"] == "pixel_farm" else 1
        self.setStyleSheet(f"""
            QWidget#WindowTitleBar {{
                background: {C_CARD};
                border: none;
                border-bottom: {border_width}px solid {C_BORDER};
            }}
            QLabel {{ background: transparent; border: none; color: {C_TEXT}; }}
            QPushButton {{
                background: transparent;
                color: {C_TEXT_SECONDARY};
                border: none;
                border-radius: {palette['radius']}px;
                padding: 0px;
                font-family: "Segoe UI Symbol", "Segoe UI";
                font-size: 17px;
                font-weight: 700;
            }}
            QPushButton:hover {{ background: {palette['hover']}; color: {C_TEXT}; }}
            QPushButton#CloseButton:hover {{ background: {C_ACCENT_RED}; color: white; }}
        """)

        layout = QHBoxLayout(self)
        layout.setContentsMargins(14, 0, 8, 0)
        layout.setSpacing(10)

        logo = QLabel()
        logo.setFixedSize(30, 30)
        icon_path = Path(__file__).resolve().parent / "mailcollector-icon.ico"
        if icon_path.exists():
            logo.setPixmap(QIcon(str(icon_path)).pixmap(26, 26))
        else:
            logo.setText("📮")
            logo.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(logo)

        brand_box = QVBoxLayout()
        brand_box.setContentsMargins(0, 5, 0, 5)
        brand_box.setSpacing(0)
        brand = QLabel("MailCollector")
        brand.setFont(QFont("Segoe UI", 11, QFont.Bold))
        brand_box.addWidget(brand)
        self.section_label = QLabel("安全邮件工作台")
        self.section_label.setFont(QFont("Microsoft YaHei UI", 8))
        self.section_label.setStyleSheet(f"color: {C_TEXT_SECONDARY};")
        brand_box.addWidget(self.section_label)
        layout.addLayout(brand_box)

        theme_badge = QLabel(current_theme()["name"])
        theme_badge.setAlignment(Qt.AlignmentFlag.AlignCenter)
        theme_badge.setStyleSheet(
            f"background: {C_SURFACE}; color: {C_PRIMARY_DARK}; border: 2px solid {C_BORDER}; "
            f"border-radius: {palette['radius']}px; padding: 4px 10px; font-size: 11px; font-weight: 800;"
        )
        layout.addWidget(theme_badge)
        layout.addStretch()

        self.minimize_button = QPushButton("−")
        self.minimize_button.setToolTip("最小化")
        self.minimize_button.clicked.connect(window.showMinimized)
        layout.addWidget(self.minimize_button)

        self.maximize_button = QPushButton("□")
        self.maximize_button.setToolTip("最大化")
        self.maximize_button.clicked.connect(self.toggle_maximize)
        layout.addWidget(self.maximize_button)

        self.close_button = QPushButton("×")
        self.close_button.setObjectName("CloseButton")
        self.close_button.setToolTip("关闭")
        self.close_button.clicked.connect(window.close)
        layout.addWidget(self.close_button)

        for button in (self.minimize_button, self.maximize_button, self.close_button):
            button.setFixedSize(42, 34)
            button.setCursor(Qt.CursorShape.PointingHandCursor)

    def set_section(self, index):
        sections = ["收集助手", "批量下载", "我的方案", "设置"]
        self.section_label.setText(translate_text(sections[index] if 0 <= index < len(sections) else "安全邮件工作台"))

    def toggle_maximize(self):
        if self.window.isMaximized():
            self.window.showNormal()
        else:
            self.window.showMaximized()
        self.update_state()

    def update_state(self):
        maximized = self.window.isMaximized()
        self.maximize_button.setText("❐" if maximized else "□")
        self.maximize_button.setToolTip(
            translate_text("还原") if maximized else translate_text("最大化")
        )

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            handle = self.window.windowHandle()
            if handle:
                handle.startSystemMove()
            event.accept()
            return
        super().mousePressEvent(event)

    def mouseDoubleClickEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.toggle_maximize()
            event.accept()
            return
        super().mouseDoubleClickEvent(event)


class WindowResizeHandle(QWidget):
    def __init__(self, window, edges, cursor):
        super().__init__(window)
        self.window = window
        self.edges = edges
        self.setCursor(cursor)
        self.setStyleSheet("background: transparent;")

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton and not self.window.isMaximized():
            handle = self.window.windowHandle()
            if handle:
                handle.startSystemResize(self.edges)
            event.accept()
            return
        super().mousePressEvent(event)


# ─── 页面3: 设置 ───

class SettingsPage(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main = main_window

        layout = QVBoxLayout(self)
        layout.setContentsMargins(40, 30, 40, 30)
        layout.setSpacing(18)

        header = QHBoxLayout()
        icon = QLabel("🎨" if APP_SETTINGS["theme"] != "pixel_farm" else "🌻")
        icon.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon.setFixedSize(58, 58)
        icon.setStyleSheet(f"background: {C_CARD}; border: 3px solid {C_BORDER}; border-radius: 12px; font-size: 28px;")
        header.addWidget(icon)
        title_box = QVBoxLayout()
        title_box.addWidget(TitleLabel("外观与语言"))
        title_box.addWidget(SubtitleLabel("让 MailCollector 更像你自己的工作台"))
        header.addLayout(title_box)
        header.addStretch()
        layout.addLayout(header)

        appearance_card = Card()
        card_layout = QVBoxLayout(appearance_card)
        card_layout.setContentsMargins(24, 22, 24, 22)
        card_layout.setSpacing(16)
        card_layout.addWidget(SectionLabel("界面主题"))

        self.theme_control = SegmentedControl([
            ("pixel_farm", "🌾 像素田园"),
            ("fresh", "🍏 清新果园"),
            ("midnight", "🌊 深海夜色"),
        ], default_value=APP_SETTINGS["theme"])
        self.theme_control.value_changed.connect(self._theme_changed)
        card_layout.addWidget(self.theme_control)

        previews = QHBoxLayout()
        previews.setSpacing(12)
        for theme_name in ("pixel_farm", "fresh", "midnight"):
            palette = THEMES[theme_name]
            preview = QFrame()
            preview.setFixedHeight(82)
            preview.setStyleSheet(
                f"QFrame {{ background: {palette['card']}; border: 3px solid {palette['border']}; "
                f"border-radius: {palette['radius']}px; }}"
            )
            preview_layout = QVBoxLayout(preview)
            preview_layout.setContentsMargins(12, 9, 12, 9)
            preview_layout.setSpacing(5)
            name = QLabel(palette["name"])
            name.setStyleSheet(f"color: {palette['text']}; font-weight: 800; background: transparent; border: none;")
            swatch = QLabel("■  ■  ■")
            swatch.setStyleSheet(
                f"color: {palette['primary']}; background: {palette['surface']}; "
                f"border: 2px solid {palette['border']}; padding: 3px 8px;"
            )
            preview_layout.addWidget(name)
            preview_layout.addWidget(swatch)
            previews.addWidget(preview)
        card_layout.addLayout(previews)

        card_layout.addWidget(SectionLabel("界面语言"))
        self.language_control = SegmentedControl([
            ("zh_CN", "🇨🇳 简体中文"),
            ("en_US", "🇬🇧 English"),
        ], default_value=APP_SETTINGS["language"])
        self.language_control.value_changed.connect(self._language_changed)
        card_layout.addWidget(self.language_control)

        hint = QLabel("主题和语言会立即生效，并在下次启动时保留。")
        hint.setWordWrap(True)
        hint.setStyleSheet(f"color: {C_TEXT_SECONDARY}; background: {C_SURFACE}; border: 2px solid {C_BORDER}; border-radius: 8px; padding: 10px;")
        card_layout.addWidget(hint)
        layout.addWidget(appearance_card)

        farm_note = Card()
        note_layout = QHBoxLayout(farm_note)
        note_layout.setContentsMargins(22, 16, 22, 16)
        note_icon = QLabel("📮")
        note_icon.setStyleSheet("font-size: 30px;")
        note_layout.addWidget(note_icon)
        note_text = QLabel("像素田园是为 MailCollector 原创设计的温暖农场风格，不包含任何游戏素材。")
        note_text.setWordWrap(True)
        note_layout.addWidget(note_text, 1)
        layout.addWidget(farm_note)
        layout.addStretch()

    def _theme_changed(self, value):
        if value != APP_SETTINGS["theme"]:
            QTimer.singleShot(0, lambda: self.main.apply_preferences(theme=value))

    def _language_changed(self, value):
        if value != APP_SETTINGS["language"]:
            QTimer.singleShot(0, lambda: self.main.apply_preferences(language=value))


# ─── 底部导航栏 ───

class NavBar(QWidget):
    def __init__(self, page_names, on_changed):
        super().__init__()

        self.setFixedHeight(74)
        self.setStyleSheet(f"""
            QWidget {{
                background: {C_BG};
                border-top: 1px solid {C_BORDER};
            }}
        """)

        layout = QHBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        self.buttons = []

        for i, (icon, name) in enumerate(page_names):
            btn = QPushButton(f"{icon}\n{name}")
            btn.setFont(QFont("Microsoft YaHei UI", 11))
            btn.setFixedHeight(74)
            btn.clicked.connect(lambda checked=False, idx=i: on_changed(idx))
            layout.addWidget(btn)
            self.buttons.append(btn)

        self.set_active(0)

    def set_active(self, index):
        for i, btn in enumerate(self.buttons):
            active = i == index
            color = C_PRIMARY if active else C_TEXT_SECONDARY
            border = C_PRIMARY if active else "transparent"
            weight = "800" if active else "600"

            btn.setStyleSheet(f"""
                QPushButton {{
                    background: transparent;
                    color: {color};
                    border: none;
                    border-top: 4px solid {border};
                    font-weight: {weight};
                    font-size: 11px;
                    border-radius: 0px;
                    padding-top: 6px;
                }}
                QPushButton:hover {{
                    color: {C_PRIMARY};
                    background: #FAFAFA;
                }}
            """)


# ─── 主窗口 ───

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()

        self.setWindowTitle("MailCollector - 邮件附件收集器")
        self.setWindowFlags(self.windowFlags() | Qt.WindowType.FramelessWindowHint)
        icon_path = Path(__file__).resolve().parent / "mailcollector-icon.ico"
        if icon_path.exists():
            self.setWindowIcon(QIcon(str(icon_path)))
        self.setMinimumSize(860, 680)
        self.resize(1040, 780)

        self._init_tray()

        # 定时任务相关
        self.scheduled_workers = []

        self.schedule_timer = QTimer(self)
        self.schedule_timer.timeout.connect(self._check_schedules)
        self.schedule_timer.start(30000)  # 每 30 秒检查一次

        self._build_ui()
        self._create_resize_handles()

    def _init_tray(self):
        icon_path = Path(__file__).resolve().parent / "mailcollector-icon.ico"
        self.tray_icon = QSystemTrayIcon(self)
        if icon_path.exists():
            self.tray_icon.setIcon(QIcon(str(icon_path)))
        else:
            self.tray_icon.setIcon(self.style().standardIcon(
                self.style().StandardPixmap.SP_MessageBoxQuestion))

        tray_menu = QMenu()
        show_action = tray_menu.addAction("显示主窗口")
        show_action.triggered.connect(self._show_from_tray)
        tray_menu.addSeparator()
        quit_action = tray_menu.addAction("退出程序")
        quit_action.triggered.connect(self._quit_app)

        self.tray_icon.setContextMenu(tray_menu)
        self.tray_icon.activated.connect(self._on_tray_activated)
        self.tray_icon.setToolTip("MailCollector - 邮件收集器")
        self.tray_icon.show()

    def _on_tray_activated(self, reason):
        if reason == QSystemTrayIcon.ActivationReason.DoubleClick:
            self._show_from_tray()

    def _show_from_tray(self):
        self.showNormal()
        self.activateWindow()
        self.raise_()

    def _quit_app(self):
        self.tray_icon.hide()
        QApplication.quit()

    def closeEvent(self, event):
        if hasattr(self, "tray_icon") and self.tray_icon and self.tray_icon.isVisible():
            event.ignore()
            self.hide()
            self.tray_icon.showMessage(
                "MailCollector 仍在运行",
                "程序已最小化到系统托盘，定时任务继续执行中。",
                QSystemTrayIcon.MessageIcon.Information,
                3000
            )
        else:
            event.accept()

    def _capture_create_state(self):
        if not hasattr(self, "create_page"):
            return None
        page = self.create_page
        try:
            imap_config = page._get_imap_config(include_password=True)
        except Exception:
            imap_config = None
        return {
            "id": page.editing_plan_id,
            "name": page.editing_plan_name,
            "conditions": page._get_conditions(),
            "action": page._get_action(),
            "output_dir": page.output_dir_edit.text(),
            "mail_source": page.mail_source.value(),
            "imap_config": imap_config,
        }

    def _restore_create_state(self, state):
        if not state:
            return
        runtime_config = dict(state.get("imap_config") or {})
        state_for_load = dict(state)
        if runtime_config:
            state_for_load["imap_config"] = {k: v for k, v in runtime_config.items() if k != "password"}
        self.create_page.load_plan_for_edit(state_for_load)
        self.create_page.editing_plan_id = state.get("id")
        self.create_page.editing_plan_name = state.get("name", "")
        if runtime_config.get("password"):
            self.create_page.imap_password.setText(runtime_config["password"])

    def _build_ui(self, active_index=0, create_state=None):
        set_theme_constants(APP_SETTINGS["theme"])
        self.setStyleSheet(build_style(APP_SETTINGS["theme"]))

        old_central = self.centralWidget()
        central = QWidget()
        self.setCentralWidget(central)
        main_layout = QVBoxLayout(central)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        self.title_bar = WindowTitleBar(self)
        main_layout.addWidget(self.title_bar)

        self.nav = QStackedWidget()
        self.create_page = CreateTaskPage(self)
        self.batch_page = BatchDownloadPage(self)
        self.plans_page = PlansPage(self)
        self.settings_page = SettingsPage(self)
        self.nav.addWidget(self.create_page)
        self.nav.addWidget(self.batch_page)
        self.nav.addWidget(self.plans_page)
        self.nav.addWidget(self.settings_page)
        main_layout.addWidget(self.nav)

        nav_icon = "🌾" if APP_SETTINGS["theme"] == "pixel_farm" else "🦉"
        self.navbar = NavBar([
            (nav_icon, "收集助手"),
            ("📦", "批量下载"),
            ("📚", "我的方案"),
            ("⚙️", "设置"),
        ], self._on_nav)
        main_layout.addWidget(self.navbar)

        self._restore_create_state(create_state)
        retranslate_widget_tree(self)
        self.nav.setCurrentIndex(active_index)
        self.navbar.set_active(active_index)
        self.title_bar.set_section(active_index)
        if old_central:
            old_central.deleteLater()
        for handle in getattr(self, "resize_handles", []):
            handle.raise_()

    def _create_resize_handles(self):
        specs = [
            (Qt.Edge.TopEdge, Qt.CursorShape.SizeVerCursor),
            (Qt.Edge.BottomEdge, Qt.CursorShape.SizeVerCursor),
            (Qt.Edge.LeftEdge, Qt.CursorShape.SizeHorCursor),
            (Qt.Edge.RightEdge, Qt.CursorShape.SizeHorCursor),
            (Qt.Edge.TopEdge | Qt.Edge.LeftEdge, Qt.CursorShape.SizeFDiagCursor),
            (Qt.Edge.BottomEdge | Qt.Edge.RightEdge, Qt.CursorShape.SizeFDiagCursor),
            (Qt.Edge.TopEdge | Qt.Edge.RightEdge, Qt.CursorShape.SizeBDiagCursor),
            (Qt.Edge.BottomEdge | Qt.Edge.LeftEdge, Qt.CursorShape.SizeBDiagCursor),
        ]
        self.resize_handles = [WindowResizeHandle(self, edges, cursor) for edges, cursor in specs]
        self._position_resize_handles()

    def _position_resize_handles(self):
        if not hasattr(self, "resize_handles"):
            return
        width, height, border, corner = self.width(), self.height(), 6, 14
        geometries = [
            (corner, 0, max(0, width - corner * 2), border),
            (corner, height - border, max(0, width - corner * 2), border),
            (0, corner, border, max(0, height - corner * 2)),
            (width - border, corner, border, max(0, height - corner * 2)),
            (0, 0, corner, corner),
            (width - corner, height - corner, corner, corner),
            (width - corner, 0, corner, corner),
            (0, height - corner, corner, corner),
        ]
        for handle, geometry in zip(self.resize_handles, geometries):
            handle.setGeometry(*geometry)
            handle.setVisible(not self.isMaximized())
            handle.raise_()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._position_resize_handles()

    def changeEvent(self, event):
        super().changeEvent(event)
        if event.type() == QEvent.Type.WindowStateChange:
            if hasattr(self, "title_bar"):
                self.title_bar.update_state()
            self._position_resize_handles()

    def apply_preferences(self, theme=None, language=None):
        global CURRENT_LANGUAGE
        create_state = self._capture_create_state()
        if theme in THEMES:
            APP_SETTINGS["theme"] = theme
        if language in ("zh_CN", "en_US"):
            APP_SETTINGS["language"] = language
            CURRENT_LANGUAGE = language
        save_app_settings(APP_SETTINGS)
        self._build_ui(active_index=2, create_state=create_state)

    def _on_nav(self, index):
        self.nav.setCurrentIndex(index)
        self.navbar.set_active(index)
        self.title_bar.set_section(index)

        if index == 2:
            self.plans_page.refresh()
            retranslate_widget_tree(self.plans_page)

    def run_plan(self, plan, scheduled=False):
        """
        统一执行方案。
        scheduled=False：手动执行，会跳回收集助手页面并使用页面进度条。
        scheduled=True：定时执行，后台执行并去重。
        """
        fresh_plan = plan
        for p in load_plans():
            if p.get("id") == plan.get("id"):
                fresh_plan = p
                break

        output = fresh_plan.get("output_dir", str(DEFAULT_OUTPUT))
        os.makedirs(output, exist_ok=True)

        conditions = refresh_dynamic_dates(fresh_plan.get("conditions", {}))
        action = fresh_plan.get("action", "both")

        if scheduled:
            schedule = fresh_plan.get("schedule", {})

            conditions["_dedupe"] = True
            conditions["_dedupe_root"] = output
            if schedule.get("mode") == ScheduleDialog.MODE_INTERVAL:
                conditions["_exclude_ids"] = set(schedule.get("collected_ids", []))

            source = fresh_plan.get("mail_source", "outlook")
            if source in ("qq", "imap"):
                imap_config = dict(fresh_plan.get("imap_config") or {})
                if imap_config:
                    imap_config["password"] = load_imap_secret(imap_config.get("account", ""))
                    if not imap_config["password"]:
                        self._on_scheduled_error(
                            "系统凭据库中没有找到该邮箱的授权码，请编辑方案并重新保存。",
                            None,
                            fresh_plan,
                        )
                        return
                    worker = IMAPWorker(conditions, output, action, imap_config)
                else:
                    self._on_scheduled_error("IMAP 方案缺少配置，已跳过", None, fresh_plan)
                    return
            else:
                worker = OutlookWorker(conditions, output, action)

            archive_current_files(
                output,
                schedule.get("archive_mode", "by_date"),
                schedule
            )
            self.scheduled_workers.append(worker)

            worker.finished.connect(
                lambda result, w=worker, p=fresh_plan: self._on_scheduled_finished(result, w, p)
            )
            worker.error.connect(
                lambda err, w=worker, p=fresh_plan: self._on_scheduled_error(err, w, p)
            )

            worker.start()

        else:
            self.create_page.load_plan_for_edit(fresh_plan)
            self.nav.setCurrentIndex(0)
            self.navbar.set_active(0)
            self.create_page._run()

    def _check_schedules(self):
        now = datetime.now()
        today_str = now.strftime("%Y-%m-%d")
        current_time = now.strftime("%H:%M")

        plans = load_plans()
        changed = False

        for plan in plans:
            schedule = plan.get("schedule", {})

            if not schedule.get("enabled"):
                continue

            end_date = schedule.get("end_date", "")
            if end_date and today_str > end_date:
                schedule["enabled"] = False
                changed = True
                continue

            mode = schedule.get("mode", ScheduleDialog.MODE_DAILY)

            if mode == ScheduleDialog.MODE_INTERVAL:
                if schedule.get("last_run", "")[:10] != today_str:
                    schedule["collected_ids"] = []
                    schedule["paused_today"] = False
                    changed = True

            if mode == ScheduleDialog.MODE_DAILY:
                if schedule.get("last_run", "")[:10] == today_str:
                    continue
                if schedule.get("time") != current_time:
                    continue

            elif mode == ScheduleDialog.MODE_INTERVAL:
                if schedule.get("paused_today"):
                    continue
                window_start = schedule.get("window_start", "12:00")
                window_end = schedule.get("window_end", "18:00")
                if current_time < window_start or current_time > window_end:
                    continue
                interval = schedule.get("interval_minutes", 10)
                last_run_str = schedule.get("last_run", "")
                if last_run_str:
                    try:
                        last_run = datetime.strptime(last_run_str, "%Y-%m-%d %H:%M:%S")
                        delta = (now - last_run).total_seconds()
                        if delta < interval * 60:
                            continue
                    except ValueError:
                        pass

            elif mode == ScheduleDialog.MODE_WEEKLY:
                if schedule.get("last_run", "")[:10] == today_str:
                    continue
                weekday = now.isoweekday()
                weekly_days = schedule.get("weekly_days", [])
                if weekday not in weekly_days:
                    continue
                weekly_time = schedule.get("weekly_time", "18:00")
                if current_time < weekly_time:
                    continue

            schedule["last_run"] = now.strftime("%Y-%m-%d %H:%M:%S")
            changed = True
            self.run_plan(plan, scheduled=True)

        if changed:
            save_plans(plans)
            self.plans_page.refresh()

    def _on_scheduled_finished(self, result, worker, plan):
        try:
            self.scheduled_workers.remove(worker)
        except ValueError:
            pass

        matched = result.get("matched", 0)
        saved = result.get("saved", 0)
        new_ids = result.get("saved_ids", [])

        print(
            f"[定时完成] {plan.get('name', '')} | "
            f"匹配 {matched} 封 | "
            f"保存 {saved} 个 | "
            f"跳过重复 {result.get('skipped_duplicate', 0)} 个"
        )

        if saved > 0 and new_ids:
            plans = load_plans()
            for p in plans:
                if p.get("id") == plan.get("id"):
                    schedule = p.get("schedule", {})
                    existing = set(schedule.get("collected_ids", []))
                    existing.update(new_ids)
                    schedule["collected_ids"] = list(existing)

                    if schedule.get("mode") == ScheduleDialog.MODE_INTERVAL and not schedule.get("pause_after_capture"):
                        reply = QMessageBox.question(
                            self,
                            "📬 检测到新邮件",
                            f"方案「{plan.get('name', '')}」捕获到 {matched} 封邮件，保存 {saved} 个文件。\n\n"
                            f"是否继续监控？\n选择「是」继续按间隔刷新到窗口结束时间。\n选择「否」则今天不再检查。",
                            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                            QMessageBox.StandardButton.Yes
                        )
                        if reply == QMessageBox.StandardButton.No:
                            schedule["paused_today"] = True

                    save_plans(plans)
                    break

            output_dir = plan.get("output_dir", str(DEFAULT_OUTPUT))
            if hasattr(self, "tray_icon") and self.tray_icon:
                self.tray_icon.showMessage(
                    f"📬 {plan.get('name', '')} 有新邮件",
                    f"保存 {saved} 个文件到：\n{output_dir}",
                    QIcon(str(Path(__file__).resolve().parent / "mailcollector-icon.ico")),
                    8000
                )

    def _on_scheduled_error(self, err, worker, plan):
        try:
            self.scheduled_workers.remove(worker)
        except ValueError:
            pass

        if hasattr(self, "tray_icon") and self.tray_icon:
            self.tray_icon.showMessage(
                f"⚠️ {plan.get('name', '')} 执行失败",
                str(err)[:200],
                QSystemTrayIcon.MessageIcon.Warning,
                8000
            )

        QMessageBox.warning(
            self,
            "定时任务失败",
            f"方案「{plan.get('name', '')}」执行失败：\n{err}"
        )


# ─── 入口 ───

def main():
    app = QApplication(sys.argv)
    app.setStyle("Fusion")

    window = MainWindow()
    window.show()

    sys.exit(app.exec())


if __name__ == "__main__":
    main()
